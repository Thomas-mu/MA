"""Audit a completed joint live trial and append documented results to the deck."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import shutil
import subprocess
import tempfile

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation

import common_comparison as common
from joint_phone_vibration_experiment import METHODS, LABELS, summarize_joint
from replay_phone_pilot import read_rows, extract_windows, phase_of
from build_phone_presentation import Deck, NAVY, MUTED, TEAL, fmt, export_pdf_and_previews


COLORS = dict(zip(METHODS, ['#007F86','#7454C4','#D67B19']))


def write_json(path, value):
    Path(path).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')


def audit(run, bundle_directory):
    meta=common.read_json(run/'run.json')
    if meta['status'] not in ('stopped','completed') or meta['cleanup_errors']:
        raise ValueError('Run did not stop cleanly.')
    for filename,digest in meta['file_sha256'].items():
        if common.sha256(run/filename)!=digest:
            raise ValueError('Original changed: '+filename)
    bundle=common.load_bundle(bundle_directory)
    if common.sha256(bundle_directory/'bundle.json')!=meta['bundle_sha256']:
        raise ValueError('Bundle mismatch')
    rows=read_rows(run/'decisions.csv');resources=read_rows(run/'resources.csv')
    cues=[json.loads(s) for s in (run/'cues.jsonl').read_text().splitlines()]
    summary=summarize_joint(rows,cues)
    by_method={m:[r for r in rows if r['method']==m] for m in METHODS}
    windows,partial=extract_windows(read_rows(run/'raw.csv'),by_method[METHODS[0]],bundle['window_size'])
    if partial!=meta['acquisition']['partial_samples_saved'] or len(windows)!=summary['shared_windows']:
        raise ValueError('Raw sample counts disagree.')
    if len(resources)!=len(windows) or len({r['window'] for r in resources})!=len(windows):
        raise ValueError('Resource/window count mismatch.')
    if meta['acquisition']['windows_formed']!=len(windows) or meta['acquisition']['windows_dropped']:
        raise ValueError('Incomplete delivery; this report requires all windows.')
    maximum=0.;compared=0
    for method in METHODS:
        function,_=common.scorer(method,bundle,bundle_directory,runtime='litert',threads=1)
        for raw,row in zip(windows,by_method[method]):
            if hashlib.sha256(raw.tobytes()).hexdigest()!=row['input_sha256']:
                raise ValueError('Stored input hash differs from raw values.')
            score,prediction=common.classify_raw(raw,bundle['scaler'],function,bundle['thresholds'][method]['value'])
            delta=abs(score-float(row['score']));maximum=max(maximum,delta)
            if prediction!=int(row['prediction']) or delta>1e-5:
                raise ValueError('Replay does not reproduce live output.')
            compared+=1
    starts=[c['monotonic_ns'] for c in cues if c['event']=='call_requested']
    ends=[c['monotonic_ns'] for c in cues if c['event']=='end_reported']
    if len(starts)!=1 or len(ends)!=1:
        raise ValueError('Expected one cue pair.')
    first_rows={m:min((r for r in rr if r['prediction']=='1' and phase_of(r,starts[0],ends[0])=='between_cues'),
                      key=lambda r:int(r['decision_ns']),default=None) for m,rr in by_method.items()}
    times=sorted({int(r['decision_ns']) for r in first_rows.values() if r is not None})
    for method,rr in by_method.items():
        entry=summary['methods'][method];first=first_rows[method]
        entry.update(processing_median_ms=float(np.median([float(r['processing_ms']) for r in rr])),
            processing_p95_ms=float(np.percentile([float(r['processing_ms']) for r in rr],95)),
            first_alarm_decision_ns=int(first['decision_ns']) if first else None,
            first_alarm_output_rank=times.index(int(first['decision_ns']))+1 if first else None,
            first_alarm_output_after_earliest_ms=(int(first['decision_ns'])-times[0])/1e6 if first else None)
    summary.update(source_run=str(run),resource_scope='combined_process_all_methods',
        verification=dict(reproduced_decisions=compared,maximum_score_difference=maximum,raw_hashes_verified=len(windows)),
        acquisition=meta['acquisition'],cue_start_s=(starts[0]-meta['start_monotonic_ns'])/1e9,
        cue_end_s=(ends[0]-meta['start_monotonic_ns'])/1e9,
        resources=dict(cpu_median_percent_one_core=float(np.median([float(r['process_cpu_percent_one_core']) for r in resources])),
            max_sampled_rss_mib=max(float(r['process_rss_bytes']) for r in resources)/2**20,
            batch_median_ms=float(np.median([float(r['batch_ms']) for r in resources])),
            cpu_frequency_khz_range=[min(float(r['cpu_frequency_khz']) for r in resources),max(float(r['cpu_frequency_khz']) for r in resources)]))
    return meta,rows,resources,cues,summary


def figures(output,meta,rows,resources,summary):
    directory=output/'assets';directory.mkdir()
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False})
    assets={}
    def save(fig,name):
        for suffix in ('png','pdf'):fig.savefig(directory/f'{name}.{suffix}',dpi=180,bbox_inches='tight',facecolor='white')
        plt.close(fig);assets[name]=directory/f'{name}.png'
    fig,axes=plt.subplots(1,2,figsize=(12.2,3.35),gridspec_kw={'width_ratios':[1,1.2]})
    earliest=min(v['first_alarm_window_between_cues'] for v in summary['methods'].values())
    for i,method in enumerate(METHODS):
        v=summary['methods'][method];color=COLORS[method]
        offset=v['first_alarm_window_between_cues']-earliest
        axes[0].scatter(offset,i,s=360,color=color,zorder=3)
        axes[0].text(offset,i,str(v['first_alarm_window_rank']),color='white',ha='center',va='center',weight='bold',fontsize=13)
        axes[0].text(offset+.18,i,f"Fenster {v['first_alarm_window_between_cues']}",va='center')
        delay=v['first_alarm_output_after_earliest_ms']
        axes[1].hlines(i,0,delay,color=color,alpha=.35,lw=4)
        axes[1].scatter(delay,i,s=360,color=color,zorder=3)
        axes[1].text(delay,i,str(v['first_alarm_output_rank']),color='white',ha='center',va='center',weight='bold',fontsize=13)
        axes[1].text(delay+34,i,fmt(delay,2)+' ms',va='center')
    for ax in axes:
        ax.set(ylim=(2.65,-.6),yticks=range(3),yticklabels=['AE','IF','RMS']);ax.grid(alpha=.18,axis='x')
    axes[0].set(title='Erstes Alarmfenster: AE und IF gleichauf',xlabel='Fenster nach frühestem Alarmfenster',xlim=(-.35,2.8),xticks=[0,1,2])
    axes[1].set(title='Erste Alarmentscheidung: AE zuerst',xlabel='Zeit nach frühester Alarmentscheidung [ms]',xlim=(-40,930))
    fig.tight_layout(w_pad=2);save(fig,'joint_alarm_order')
    fig,axes=plt.subplots(3,1,figsize=(12.2,4.7),sharex=True)
    for ax,method in zip(axes,METHODS):
        selected=[r for r in rows if r['method']==method]
        x=[(int(r['window_complete_ns'])-meta['start_monotonic_ns'])/1e9 for r in selected]
        ax.plot(x,[float(r['score'])/float(r['threshold']) for r in selected],color=COLORS[method],lw=1.7)
        ax.axhline(1,color='#AB5E06',linestyle='--',lw=1)
        ax.axvspan(summary['cue_start_s'],summary['cue_end_s'],color='#CBD9E4',alpha=.35)
        ax.set_ylabel(LABELS[method]+'\nScore / Schwelle',fontsize=10);ax.set_ylim(bottom=0);ax.grid(alpha=.2)
    axes[-1].set_xlabel('Gemeinsame Fenster-Endzeit seit Aufnahmestart [s]')
    fig.tight_layout(h_pad=.5);save(fig,'joint_scores')
    fig,axes=plt.subplots(1,2,figsize=(12.2,3.3))
    x=[float(r['elapsed_s']) for r in resources]
    axes[0].plot(x,[float(r['process_cpu_percent_one_core']) for r in resources],color='#007F86')
    axes[0].set(title='CPU des gesamten Aufnahmeprozesses',ylabel='% eines Kerns',xlabel='Zeit seit Start [s]',ylim=(0,None))
    axes[1].plot(x,[float(r['process_rss_bytes'])/2**20 for r in resources],color='#7454C4')
    axes[1].set(title='RAM des gesamten Aufnahmeprozesses',ylabel='RSS [MiB]',xlabel='Zeit seit Start [s]',ylim=(0,200))
    for ax in axes:ax.grid(alpha=.18)
    fig.tight_layout(w_pad=2);save(fig,'joint_resources')
    return assets


def report_and_screenshot(output,meta,summary):
    lines=['# Gemeinsamer Live-Anruf: Ergebnis', '',
        'Alle drei Methoden liefen live auf denselben Rohfenstern aus einem Sensorstream. '
        '75 % PWM, eingefrorene Modelle und Schwellen, kein erneutes Training. '
        'Ein Anruf, keine unabhängige physische Anfangs-/Endreferenz.', '',
        '| Methode | Erstes Alarmfenster | Fensterrang | Ausgaberang | Abstand zur frühesten Alarmentscheidung [ms] | Alarmfenster |',
        '| --- | ---: | ---: | ---: | ---: | ---: |']
    html_rows=[]
    for method in METHODS:
        v=summary['methods'][method]
        lines.append(f"| {LABELS[method]} | {v['first_alarm_window_between_cues']} | {v['first_alarm_window_rank']} | {v['first_alarm_output_rank']} | {fmt(v['first_alarm_output_after_earliest_ms'],3)} | {v['alarm_windows']} |")
        html_rows.append(f"<tr><td>{LABELS[method]}</td><td>{v['first_alarm_window_between_cues']}</td><td>{v['alarm_windows']}</td></tr>")
    lines+=['', 'Autoencoder und Isolation Forest überschreiten die Schwelle im selben Fenster. '
        'Der Autoencoder liefert seine Entscheidung rund 19,06 ms früher. RMS überschreitet die Schwelle ein Fenster später '
        'und liefert seine erste Alarmentscheidung rund 615,66 ms nach dem Autoencoder. '
        'Das sind beobachtete relative Ausgabeabstände dieses gemeinsamen Laufs, keine absolute Fehler-Erkennungsverzögerung '
        'ab physischem Vibrationsbeginn. Thread-Scheduling und konkurrierende Berechnung beeinflussen die Ausgabezeiten.', '',
        '![Nummerierte Rangfolge](assets/joint_alarm_order.png)', '',
        '## Gemeinsame Daten und Phasen', '',
        f"{summary['shared_windows']} identische vollständige Fenster, {summary['verification']['reproduced_decisions']} gültige Entscheidungen. "
        f"{summary['acquisition']['samples_captured']} Rohsamples, davon {summary['acquisition']['partial_samples_saved']} in einem unvollständigen Schlussfenster. "
        'Keine verworfenen Fenster, keine gemeldeten Sensorlücken/Überläufe/Sättigungen. '
        'Rohwerte, Eingabehashes, Fenstergrenzen und alle Entscheidungen wurden nachgerechnet; maximale Score-Abweichung 0.', '',
        'Vollständig vor Freigabe: 77 Fenster ohne Alarm bei allen Methoden. '
        'Zwischen Freigabe und Rückmeldung: 79 Fenster, davon AE 11 / IF 7 / RMS 6 Alarmfenster. '
        'Nach Rückmeldung: 39 Fenster ohne Alarm. Zwei grenzüberlappende Fenster sind in dieser Aufteilung ausgeschlossen, '
        'in der Gesamtsumme enthalten und ebenfalls ohne Alarm. Die Chat-Abschnitte sind nicht die physische Vibrationsdauer.', '',
        '![Gemeinsame Scores](assets/joint_scores.png)', '', '## Einfluss auf den Pi', '']
    r=summary['resources']
    lines += [f"Prozess-CPU-Median: {fmt(r['cpu_median_percent_one_core'],2)} % eines Kerns; "
        f"maximal abgetastete RSS: {fmt(r['max_sampled_rss_mib'],2)} MiB. "
        'Das sind Werte des gesamten gemeinsamen Prozesses einschließlich Sensorerfassung, drei Methoden und Logging; '
        'keine getrennte CPU-/RAM-Zuordnung zu den Methoden.', '',
        '| Methode | Berechnungszeit Median [ms] | P95 [ms] |', '| --- | ---: | ---: |']
    for method in METHODS:
        v=summary['methods'][method]
        lines.append(f"| {LABELS[method]} | {fmt(v['processing_median_ms'],3)} | {fmt(v['processing_p95_ms'],3)} |")
    lines += ['', 'Berechnungszeiten enthalten Standardisierung und Entscheidung innerhalb des jeweiligen Worker-Threads. '
        'Sie entstehen unter Konkurrenz im gemeinsamen Prozess, nicht in einem isolierten Benchmark. '
        'Eine Barriere gibt die Worker pro Fenster frei, garantiert aber keinen exakt gleichen CPU-Startzeitpunkt.', '',
        '![Gemeinsame Prozesslast](assets/joint_resources.png)', '',
        '## Grenzen und Dokumentation', '',
        'Ein gemeinsamer Anruf reicht nicht für eine allgemeine Sieger-Rangfolge oder Erkennungsquote. '
        'Mehr Alarmfenster bedeuten nicht mehr unabhängig erkannte Fehler. '
        'Finale Handyaufstellung und Normalkalibrierung bleiben unbestätigt; externe Tischvibration ist kein belegter Lüfterdefekt. '
        'Alle Quellen und Neustarts früherer Einzelversuche bleiben erhalten. '
        'Der Screenshot zeigt den tatsächlich im Browser gerenderten gespeicherten Bericht, keine Live-GUI oder Aufbauaufnahme.', '',
        '[Browser-Screenshot](joint_report_screenshot.png)', '',
        'Die Präsentation enthält in Teil A (Folien 1–10) die bisherigen Einzelversuche/Replay-Ergebnisse und in '
        'Teil B (Folien 11–14) diesen zusätzlichen gemeinsamen Live-Anruf.', '']
    (output/'joint_report.md').write_text('\n'.join(lines),encoding='utf-8')
    page='''<!doctype html><html lang="de"><meta charset="utf-8"><title>Gemeinsamer Live-Anruf – gespeicherter Bericht</title>
<style>*{box-sizing:border-box}body{margin:0;padding:30px 42px;background:#F5F8FA;color:#142D42;font:18px "DejaVu Sans",sans-serif}
h1{font-size:33px;margin:10px 0}h2{font-size:23px}small{color:#587084}.grid{display:grid;grid-template-columns:1fr 1fr;gap:24px;margin-top:22px}
.card{background:white;padding:22px;border-radius:12px}table{border-collapse:collapse;width:100%}td,th{text-align:left;padding:10px;border-bottom:1px solid #dae2ea}
p{line-height:1.45}.figure{background:white;margin-top:24px;border-radius:12px;padding:12px 20px}img{width:100%}.note{color:#87500F;font-size:17px}
footer{margin-top:17px;font-size:13px;color:#587084}</style>
<small>RASPBERRY PI 5 · GEMEINSAMER LIVE-VERSUCH · GESPEICHERTE AUSWERTUNG</small>
<h1>Gleiche Sensordaten: Autoencoder, Isolation Forest und RMS</h1>
<div class="grid"><div class="card"><h2>Ein Anruf, 197 gemeinsame Fenster</h2><table><tr><th>Methode</th><th>Erster Alarm</th><th>Alarmfenster</th></tr>'''
    page+=''.join(html_rows)+'''</table><p>591 gültige Entscheidungen · 0 verworfene Fenster<br>Gleiche Rohwerte und Fenstergrenzen nachgeprüft.</p></div>
<div class="card"><h2>Fenstererkennung ≠ Ausgabezeit</h2><p>AE und IF: erstes Alarmfenster 108.<br>RMS: erstes Alarmfenster 109.</p>
<p>AE-Ausgabe rund 19 ms vor IF.<br>RMS-Ausgabe rund 616 ms nach AE.</p><p class="note">Relative Ausgabeabstände dieses Laufs, keine gemessene Verzögerung ab physischem Vibrationsbeginn.</p></div></div>
<div class="figure"><img src="assets/joint_alarm_order.png" alt="Fensterrang und Rang der ersten Alarmentscheidung"></div>
<footer>Echter Browserbericht nach Messende · keine Live-GUI · kein Foto des Aufbaus · PWM 75 % · Original: joint_trials/all_methods_01</footer></html>'''
    report=output/'joint_report.html';report.write_text(page,encoding='utf-8')
    screenshot=output/'joint_report_screenshot.png'
    with tempfile.TemporaryDirectory(prefix='joint-pilot-report-') as profile:
        command=['chromium','--headless','--no-sandbox','--disable-gpu','--disable-background-networking','--no-first-run',
            f'--user-data-dir={profile}','--hide-scrollbars','--window-size=1600,1080','--virtual-time-budget=1500',
            f'--screenshot={screenshot}',report.as_uri()]
        result=subprocess.run(command,capture_output=True,text=True,timeout=45)
        (output/'browser_capture.log').write_text(result.stdout+result.stderr);result.check_returncode()
    write_json(output/'browser_capture.json',dict(captured_utc=datetime.now(timezone.utc).isoformat(),
        type='actual_browser_screenshot_of_saved_report',live_gui=False,physical_setup_photo=False,
        command=command,html_sha256=common.sha256(report),screenshot_sha256=common.sha256(screenshot)))
    return screenshot


def append_deck(output,base,assets,screenshot,summary):
    d=Deck();d.prs=Presentation(base)
    if len(d.prs.slides)!=10:raise ValueError('Expected the archived ten-slide result deck.')
    for shape in d.prs.slides[0].shapes:
        if shape.has_text_frame and shape.text=='ERGEBNISÜBERBLICK':
            shape.text_frame.paragraphs[0].runs[0].text='TEIL A · EINZELVERSUCHE UND OFFLINE-REPLAY'
    d.prs.core_properties.title='Handyvibration: Einzelversuche, Replay und gemeinsamer Live-Anruf'
    def slide(title,subtitle,notes):
        s=d.slide('Teil B · Gemeinsamer Live-Anruf',title,subtitle,notes)
        for sh in s.shapes:
            if sh.has_text_frame and sh.text.startswith('VORBEREITUNGSSTAND'):
                sh.text_frame.paragraphs[0].runs[0].text='GEMEINSAMER LIVE-PILOT · 19.09.2026 · Ein Anruf, identische Fenster · keine physische Zeitreferenz'
        return s
    s=slide('Gemeinsamer Anruf: AE liefert zuerst Alarm',
        'Gleiche Daten und dasselbe Vibrationsereignis. Fenstererkennung und Ausgabezeit sind zwei verschiedene Vergleiche.',
        'Ein zusätzlicher gemeinsamer Live-Anruf mit allen drei Methoden auf einem Sensorstream. '
        'Drei Worker-Threads erhalten identische Fenster und werden mit einer Barriere freigegeben. '
        'Erstes Alarmfenster AE und IF 108, RMS 109. Fensterrang AE/IF gemeinsam 1, RMS 2. '
        'Erste tatsächliche decision_ns: AE zuerst, IF 19,061107 ms später, RMS 615,658777 ms später. '
        'Daher Ausgabe-Ränge 1/2/3. Diese relativen Abstände enthalten Rechenzeit und Scheduling, '
        'aber sind keine absolute Verzögerung ab physischem Vibrationsbeginn. '
        'Quelle: sources/joint_run/decisions.csv; joint_summary.json.')
    d.picture(s,assets['joint_alarm_order'],.6,2.6,12.1,3.5)
    d.text(s,.8,6.25,11.8,.6,'Gleiches Alarmfenster heißt nicht gleiche Ausgabezeit. Eine allgemeine Rangfolge ist damit nicht belegt.',18,TEAL,True)
    s=slide('197 Fenster, drei Methoden – gleiche Daten',
        'Ein Sensorstream, unveränderte Schwellen. Grau: Chat-Freigabe bis Rückmeldung, keine physische Vibrationsdauer.',
        '197 vollständige Fenster zu je 128 × 3 Werten, 25.270 Rohsamples, 54 Samples im unvollständigen Schlussfenster. '
        'Keine verworfenen oder ungültigen Fenster, keine gemeldeten Lücken/Überläufe/Sättigungen. '
        '591 Entscheidungen wurden offline exakt reproduziert; Rohfensterhashes und gemeinsame Zeitgrenzen geprüft. '
        'AE 11, IF 7, RMS 6 Alarmfenster, alle vollständig zwischen den Chat-Markierungen. '
        '77 Fenster vollständig vor Freigabe, 79 zwischen Markierungen, 39 danach, 2 grenzüberlappend. '
        'Vorher/nachher und grenzüberlappend keine Alarme. Ein Anruf bleibt ein Versuch, nicht 24 unabhängige Fehler. '
        'Score/Schwelle ist keine Wahrscheinlichkeit. Quelle: raw.csv, decisions.csv, cues.jsonl, run.json.')
    d.picture(s,assets['joint_scores'],.6,2.38,12.1,4.12)
    d.text(s,.8,6.64,11.8,.3,'Alarmfenster: Autoencoder 11 · Isolation Forest 7 · RMS 6. Keine allgemeine Erkennungs- oder Fehlalarmrate.',13,MUTED)
    resource=summary['resources']
    s=slide('Die Gesamtlast im gemeinsamen Betrieb',
        'CPU/RAM enthalten alle Methoden, Sensorerfassung und Logging. Keine getrennte Modelllast aus diesen Werten ableiten.',
        'Prozess-CPU als Median intervallweiser Samples: '+fmt(resource['cpu_median_percent_one_core'],2)+' % eines Kerns; '
        'maximal abgetastete RSS '+fmt(resource['max_sampled_rss_mib'],2)+' MiB. Kein Start-Peak und keine elektrische Leistung in Watt. '
        'Worker-Berechnungszeiten enthalten Standardisierung, Score und Entscheidung unter Konkurrenz. '
        'Eine Thread-Barriere garantiert keine exakt gleichzeitige CPU-Ausführung. '
        'Betriebssystem-Scheduling, Bibliotheken und dynamischer Takt beeinflussen die Werte. '
        'Mit isolierten früheren Läufen nur explorativ vergleichen, kein kontrollierter Parallelisierungsbenchmark. '
        'Quelle: sources/joint_run/resources.csv und decisions.csv.')
    d.picture(s,assets['joint_resources'],.6,2.62,12.1,3.16)
    d.text(s,.8,6.0,11.7,.8,
        f"CPU-Median: {fmt(resource['cpu_median_percent_one_core'],1)} % eines Kerns · max. abgetastete RSS: {fmt(resource['max_sampled_rss_mib'],1)} MiB\n"
        'Berechnung Median: AE '+fmt(summary['methods']['tflite_autoencoder']['processing_median_ms'],3)+
        ' ms · IF '+fmt(summary['methods']['isolation_forest']['processing_median_ms'],3)+
        ' ms · RMS '+fmt(summary['methods']['rms']['processing_median_ms'],3)+' ms',17)
    s=slide('Gemeinsame Aufnahme vollständig dokumentiert',
        'Echter Screenshot des gespeicherten Berichts; bestehende Einzelversuche und Präsentationsversionen bleiben erhalten.',
        'Der abgebildete Bericht wurde tatsächlich im Browser gerendert, nicht als Live-GUI rekonstruiert. '
        'Raw-CSV, Eingabehashes, Zeitstempel, Entscheidungs-CSV, Ressourcen und Quellenmanifest sind archiviert. '
        'Acht fokussierte Tests sind bestanden; zusätzlich 591 echte Entscheidungen aus dem gemeinsamen Lauf exakt nachgerechnet. '
        'Trotz identischer Inputs bleibt dies ein einzelner Anruf ohne unabhängige physische Start-/Endmessung. '
        'Finaler Handyaufbau ist nicht verifiziert; Tischvibration ist kein belegter Lüfterdefekt. '
        'Schwellen und PWM 75 % unverändert. Für allgemeine Aussagen Aufbau validieren und Ereignisse wiederholen. '
        'Teil A zeigt frühere separate Anrufe mit Offline-Replay; Teil B ist der zusätzliche gemeinsame Live-Anruf.')
    d.picture(s,screenshot,.55,2.4,8.0,4.35)
    d.box(s,8.85,2.7,3.9,3.62)
    d.text(s,9.08,3.0,3.45,.43,'Für die Einordnung',21,NAVY,True)
    d.text(s,9.08,3.68,3.43,2.28,'Gleiche Daten: nachgeprüft\nEin gemeinsamer Anruf\nKeine physische Zeitreferenz\nKeine allgemeine Gütequote\nSchwellen unverändert',16)
    pptx=output/'handyvibration_mit_gemeinsamem_liveversuch.pptx';d.prs.save(pptx)
    prior=(base.parent/'speaker_notes.md').read_text() if (base.parent/'speaker_notes.md').exists() else ''
    notes=prior+'\n\n# Teil B: Gemeinsamer Live-Versuch\n\n'+'\n\n'.join(
        f'## {i}. {title}\n\n{text}' for i,(title,text) in enumerate(d.notes,11))
    (output/'speaker_notes.md').write_text(notes,encoding='utf-8')
    return pptx


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,required=True);parser.add_argument('--run',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();root=args.root.resolve();run=args.run.resolve();output=args.output.resolve()
    meta,rows,resources,cues,summary=audit(run,root/'bundle')
    # The explanatory text describes this archived pilot, not arbitrary future trials.
    expected=[(108,11,0.),(108,7,19.061107),(109,6,615.658777)]
    if summary['shared_windows']!=197 or summary['verification']['maximum_score_difference']!=0:
        raise ValueError('Narrative is specific to the verified 197-window joint pilot.')
    for method,(first,count,delay) in zip(METHODS,expected):
        entry=summary['methods'][method]
        if (entry['first_alarm_window_between_cues']!=first or entry['alarm_windows']!=count
                or not np.isclose(entry['first_alarm_output_after_earliest_ms'],delay,rtol=0,atol=1e-6)):
            raise ValueError('Update the narrative before rendering a different joint pilot.')
    output.mkdir(parents=True,exist_ok=False)
    sources=[]
    originals=[(p,Path('joint_run')/p.name) for p in run.iterdir() if p.is_file()]
    originals += [(p,Path('bundle')/p.name) for p in (root/'bundle').iterdir() if p.is_file()]
    base=root/'presentation_v2_results/handyvibration_ergebnisse.pptx'
    originals += [(base,Path('prior_deck')/base.name),(base.parent/'speaker_notes.md',Path('prior_deck/speaker_notes.md'))]
    for filename in ('build_joint_phone_results.py','build_phone_presentation.py','joint_phone_vibration_experiment.py'):
        originals.append((Path(__file__).parent/filename,Path(filename)))
    for original,relative in originals:
        target=output/'sources'/relative;target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(original,target)
        digest=common.sha256(original)
        if common.sha256(target)!=digest:raise ValueError('Source copy mismatch')
        sources.append(dict(original=str(original),snapshot=str(target.relative_to(output)),sha256=digest))
    write_json(output/'source_manifest.json',dict(created_utc=datetime.now(timezone.utc).isoformat(),sources=sources))
    write_json(output/'joint_summary.json',summary)
    assets=figures(output,meta,rows,resources,summary)
    screenshot=report_and_screenshot(output,meta,summary)
    pptx=append_deck(output,output/'sources/prior_deck'/base.name,assets,screenshot,summary)
    export_pdf_and_previews(pptx,output)
    write_json(output/'artifact_manifest.json',dict(created_utc=datetime.now(timezone.utc).isoformat(),
        files=[dict(path=str(p.relative_to(output)),sha256=common.sha256(p)) for p in sorted(output.rglob('*')) if p.is_file()]))
    print(json.dumps(dict(pptx=str(pptx),pdf=str(pptx.with_suffix('.pdf')),summary=summary),indent=2))


if __name__=='__main__':
    main()
