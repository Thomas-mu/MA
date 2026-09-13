"""Revise the existing manuscript into an explicitly bounded empirical pilot.

Reads a preserved Word backup and archived results only. No acquisition,
model evaluation, training or hardware access is performed.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, importlib.util, json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

B=Path(__file__).resolve().parent
ROOT=B.parents[1]
OLD=ROOT/'results/upright_v4_finalization_20260912_103145'

def sha(p):
    with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def main():
    state=json.loads((B/'initial_state.json').read_text())
    source=Path(state['backup']);assert sha(source)==state['input_word_sha256']
    doc=Document(source);changes=[];additions=[]
    spec=importlib.util.spec_from_file_location('preserved_style_helpers',OLD/'build_thesis.py')
    helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
    def match(prefix):
        found=[p for p in doc.paragraphs if p.text.startswith(prefix)]
        assert len(found)==1,(prefix,len(found));return found[0]
    def replace(prefix,new,reason):
        p=match(prefix);old=p.text;p.text=new
        changes.append({'kind':'paragraph','before':old,'after':new,'reason':reason})
    def insert_before(anchor,text,style=None):
        p=anchor.insert_paragraph_before(text,style=style)
        additions.append({'kind':'paragraph','text':text,'anchor':anchor.text})
        return p
    def insert_after(prefix,text):
        prior=match(prefix)
        p=doc.add_paragraph(text);prior._p.addnext(p._p)
        additions.append({'kind':'paragraph','text':text,'after':prior.text})
        return p

    # Preserve the personal declaration and its unsigned lines. Add an abstract
    # on a separate Roman-numbered page using the existing front-matter style.
    toc=match('Inhaltsverzeichnis')
    title=insert_before(toc,'Zusammenfassung','Heading 1');helper.set_num(title,num=0)
    title.paragraph_format.page_break_before=True
    abstracts=[
        'Diese Arbeit untersucht einen lokalen Methodenvergleich zur normalitätsbasierten Anomalieerkennung an einem Lüfterprüfstand. Ein ADXL345 erfasst dreiachsige Beschleunigungen; ein Raspberry Pi 5 verarbeitet identische, nicht überlappende Fenster mit 128 XYZ-Punkten. Verglichen werden eine RMS-Schwelle auf standardisierten Werten, Isolation Forest und ein TFLite-Autoencoder. Vollständige Aufnahmen sind Training, Validierung oder Test zugeordnet. Modelle, Skalierung und Schwellen bleiben in den unabhängigen Tests unverändert.',
        'Die kontrollierte Testfolge umfasst Normalbetrieb, eine äußere Luftstromveränderung durch eine separat befestigte Platte und eine normale Rückkehrreferenz bei 75 % PWM. Bewertet wird jeweils der vorab festgelegte Abschnitt [180,300) s einer 300-s-Aufnahme. Von 194 gültigen Fenstern des veränderten Zustands markieren RMS sieben, Isolation Forest vier und der Autoencoder keines. Die F1-Werte der vollständigen Folge betragen 6,93 %, 4,04 % und 0 %. Die erwartete Überlegenheit des Autoencoders ist in dieser Folge nicht bestätigt. Zusätzliche Normaltests bei 50 % PWM zeigen methoden- und aufnahmeabhängige Fehlalarme.',
        'In der getrennten Laufzeitprüfung liegt das empirische P99 der gesamten Verarbeitungslatenz im Sensorbetrieb je nach Methode und Lauf zwischen 6,113 und 28,577 ms; das beobachtete Fensterintervall beträgt etwa 618 ms. Die Studie belegt damit die lokale Ausführbarkeit im gemessenen Umfang, jedoch keine zuverlässige Erkennung der untersuchten Luftstromänderung. Eine einzige unabhängige veränderte Aufnahme, die ungemessene Drehzahl und die nicht unabhängig geklärte Abweichung zwischen nominell 200 Hz und etwa 207 XYZ/s begrenzen die Aussagekraft. Der Abschluss erfolgt als Pilotstudie mit erhaltenen negativen Ergebnissen und ausdrücklich offenen Nachweisen.'
    ]
    for text in abstracts:insert_before(toc,text)

    replace('Die Arbeitshypothesen legen die Vergleichsgrößen fest.',
        'Die Arbeitshypothesen legen die Vergleichsgrößen fest. Bereits zur Entwicklung verwendete Daten werden explorativ ausgewertet. Die unabhängige Prüfung verwendet davon getrennte Testaufnahmen und ein vor ihrer Auswertung festgelegtes Versuchsprotokoll. Ihre Aussagekraft bleibt durch Zahl und Vielfalt der vollständigen Aufnahmen begrenzt.',
        'Distinguish completed independent tests from unperformed further replication.')
    replace('Kapitel 4 übersetzt diesen Untersuchungsbedarf in überprüfbare Anforderungen',
        'Kapitel 4 und 5 konkretisieren Anforderungen und Systemauswahl.',
        'Shorten redundant chapter transition to prevent a nearly empty continuation page.')
    replace('Kapitel 6 dokumentiert Aufbau, Software, Eignungsprüfungen',
        'Kapitel 6 dokumentiert Implementierung und Testprotokoll; Kapitel 7 bewertet die unabhängigen Tests anhand unveränderter Anforderungen und Hypothesen.',
        'Shorten redundant transition; preserve test separation and unchanged assessment criteria.')
    insert_after('Dieses Kapitel begründet die Systemauswahl anhand der Anforderungen',
        'Die folgende Darstellung rekonstruiert die Systementscheidung und die dabei festgelegten Prüfvorbehalte. Die Anforderungen beschreiben den ursprünglichen Zielmaßstab. Ob die jeweiligen Nachweise im vorhandenen Datenstand erbracht wurden, bewertet Kapitel 7.')
    replace('Als konkrete Variante wird der ARCTIC P12 Pro PST vorgesehen',
        'Als konkrete Variante wurde der ARCTIC P12 Pro PST vorgesehen, dessen Datenblatt PWM-Ansteuerung und ein Drehzahlsignal ausweist (ARCTIC, o. J.). Ein PWM-Tastgrad ist eine Stellgröße und kein Nachweis einer konstanten Drehzahl. Die Konzeption sah deshalb die Erfassung der tatsächlichen Drehzahl vor. Dieser Nachweis wurde in der abgeschlossenen Pilotstudie nicht erbracht; die Betriebspunkte werden entsprechend durch ihre dokumentierten PWM-Vorgaben beschrieben. Für FF4 wurde ein zweiter normaler Betriebspunkt bei unveränderter Montage untersucht. Versuchszustand, Schutzmaßnahmen und Messablauf sind in Kapitel 6 dokumentiert; den begrenzten Nachweis der Rücksetzbarkeit bewertet Kapitel 7.',
        'Retain original RPM requirement while declaring the actual unmeasured status.')
    replace('Ein exklusiver Prozess-Lock koordiniert Programme',
        'Ein exklusiver Prozess-Lock koordiniert Programme, die diesen Steuerweg verwenden. Vor Hardwareläufen werden zusätzlich mögliche konkurrierende Steuer- und Sensorprozesse geprüft. Fremde Programme können diese Sperre umgehen; der Lock ist deshalb kein elektrischer Verriegelungsmechanismus. Die Messläufe halten die vorgegebene PWM konstant. Modellentscheidungen verändern diese Stellvorgabe während des Methodenvergleichs nicht; eine Rückwirkung über die Ansteuerung ist damit ausgeschlossen. Ob die tatsächliche Drehzahl und weitere physikalische Bedingungen zwischen getrennten Läufen identisch waren, ist dadurch nicht nachgewiesen.',
        'Constant commanded PWM does not prove identical physical states across runs.')
    insert_after('Bei den Normalreferenzen sind die auffälligen Entscheidungen Fehlalarme.',
        'Für die vollständige Folge wird airflow_modified als vorab festgelegtes positives Zustandslabel ausgewertet. Auf denselben 582 gültigen Fenstern ergeben sich F1-Werte von 6,93 % für RMS, 4,04 % für Isolation Forest und 0 % für den Autoencoder. Die F1-Differenz des Autoencoders zur statistischen Baseline beträgt somit −6,93 Prozentpunkte; H1 ist in dieser Folge nicht bestätigt. Die vollständigen Konfusionszahlen und Kennzahlen stehen im Anhang. Die 194 positiven Fenster stammen aus einer einzigen Aufnahme und belegen keine Wiederholbarkeit des Ergebnisses.')
    replace('Der Gesamtfehler ist das arithmetische Mittel der drei Achsenfehler.',
        'Der mittlere Gesamtfehler ist das arithmetische Mittel der drei mittleren Achsenfehler. Die kleineren X- und Y-Fehler überkompensieren den mit Platte geringfügig höheren Z-Fehler. Zusätzlich zeigt die Prüfung der einzelnen Modellfenster, dass in dieser Aufnahme kein Gesamtfehler die eingefrorene Schwelle von rund 0,54436 überschreitet. Ein kleinerer Gruppenmittelwert allein würde diesen Nullbefund nicht belegen. Ein größerer physischer Vektor-RMS muss bei dieser Vorverarbeitung und diesem Modell keinen größeren Anomaliescore verursachen. Auch der RMS-Modellscore wird aus standardisierten Werten berechnet und ist nicht identisch mit dem physischen Vektor-AC-RMS.',
        'Separate comparison of group means from the independently verified absence of individual threshold crossings.')
    replace('Vierzig der 51 Autoencoder-Fehlalarme in normal_before',
        'Vierzig der 51 Autoencoder-Fehlalarme in normal_before liegen zwischen 180 und 240 s ab dem Stellbefehl. Über alle 194 Modellfenster dieser Aufnahme beträgt die explorative Pearson-Korrelation zwischen Autoencoder-Score und physischer X-Achsen-RMS ungefähr r = 0,821. Die Fenster sind zeitlich abhängig; der Koeffizient beschreibt diesen Verlauf und ist kein unabhängiger Kausalnachweis. Mit der Z-Achsen-RMS beträgt die entsprechende Korrelation nur rund 0,002. Die Fehlalarme kennzeichnen daher keine einfache Überschreitung der Gesamtvibration.',
        'Specify Pearson correlation, full 194-window population and dependence from the archived exploratory analysis.')

    insert_after('Der Kernversuch bewertet die lokale Inferenz',
        'Die Abschlussfassung verwendet die bis zum 12. September 2026 vorhandenen Messungen; die abschließende Auswertung und Redaktion erfolgen am 13. September 2026. Weitere Plattenwiederholungen und unabhängige Drehzahl- oder Sensortaktmessungen werden in dieser Fassung nicht als durchgeführt behandelt. Diese Begrenzung des tatsächlich erbrachten Umfangs wird nach der Versuchsauswertung offengelegt. Die ursprünglichen Anforderungen und Hypothesen bleiben als Bewertungsmaßstab erhalten; insbesondere werden Teilnachweise der Messqualität und Wiederholbarkeit nicht nachträglich zu vollständiger Erfüllung erklärt.')
    replace('Die Softwareprüfung des abschließenden Livewegs umfasst 18 gezielte Tests.',
        '18 Softwaretests prüfen Referenzscores, Datenqualität und Fehlerpfade. Die Hardwareläufe belegen die lokale Verarbeitung im gemessenen Umfang. Kapitel 7 bewertet die Anforderungen einschließlich offener Drehzahl-, Messketten- und Wiederholungsnachweise.',
        'Remove residual contradiction about a required proven defect class; retain actual missing evidence.')
    insert_after('Das eingefrorene Paket run_001 enthält',
        'Dieselbe normale Validierungsaufnahme dient sowohl zur Auswahl der gespeicherten Autoencoder-Epoche als auch zur Schwellenkalibrierung. Sie ist damit keine von der Modellauswahl unabhängige zusätzliche Kalibrierungsstichprobe. Die unabhängigen Testaufnahmen bleiben davon getrennt. Zwei Trainingsstarts und ein Validierungsstart erfassen nur einen kleinen Teil der möglichen Normalvariation; die große Zahl benachbarter Fenster ersetzt keine zusätzliche unabhängige Aufnahme.')
    insert_after('Die RMS-Baseline ist dimensionslos.',
        'Beim Isolation Forest wird der negierte Wert von score_samples mit der gespeicherten gemeinsamen Quantilschwelle verglichen. Die bibliotheksinterne Einstellung contamination = auto wird nicht als Fehlalarmvorgabe verwendet; die Entscheidung folgt nicht der unveränderten predict-Grenze der Bibliothek. Damit bleibt die tatsächlich ausgewertete Scoredefinition von einer angenommenen Ausreißerquote getrennt.')
    replace('Bei RMS und Isolation Forest sinkt der zusammengefasste Fehlalarmanteil',
        'Bei RMS und Isolation Forest sinkt der zusammengefasste Fehlalarmanteil im vorliegenden Vergleich. Beim Autoencoder steigt er von 5,50 % auf 6,01 %, also um ungefähr 0,52 Prozentpunkte. Die sechs 75-%-Normalaufnahmen reichen beim Autoencoder jedoch von 0 bis 26,29 % Fehlalarmen; innerhalb der drei 50-%-Starts sinkt der Anteil von 8,76 auf 3,61 %. Die kleine gepoolte Zunahme beschreibt somit keine über alle Starts einheitliche Verschiebung und hängt von der festgelegten Zusammensetzung der Normalreferenzen ab. Die zeitlich getrennten Reihen isolieren die PWM nicht von weiteren Einflüssen. H3 ist für den Autoencoder nur mit der Richtung des zusammengefassten Anteils vereinbar und weder allgemein noch kausal bestätigt; für RMS und Isolation Forest zeigt der zusammengefasste Vergleich die entgegengesetzte Richtung.',
        'Clarify H3 pooling, run variability and causal boundary without changing reference selection.')

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells)==3 and row.cells[0].text=='A7':
                before=[c.text for c in row.cells]
                row.cells[1].text='Im Messumfang nachgewiesen'
                row.cells[2].text='P99 der Verarbeitung liegt in allen geprüften Läufen unter dem beobachteten Fensterintervall; keine gezählte Fristüberschreitung. Hostseitige Grenze; keine harte Echtzeitgarantie.'
                changes.append({'kind':'requirement_row','before':before,'after':[c.text for c in row.cells],'reason':'State the evidenced H2 result explicitly rather than only referring to a section.'});helper.style_table(table)
            if len(row.cells)==3 and row.cells[0].text=='A10':
                before=[c.text for c in row.cells]
                row.cells[1].text='Im festgelegten Vergleich nachgewiesen'
                row.cells[2].text='Drei normale 50-%-Tests bei unverändertem Paket; Vergleich mit den vorab bestimmten 75-%-Normalreferenzen. Keine randomisierte Paarung und kein isolierter kausaler PWM-Effekt.'
                changes.append({'kind':'requirement_row','before':before,'after':[c.text for c in row.cells],'reason':'Separate performed comparison from general or causal inference.'});helper.style_table(table)
            if len(row.cells)==3 and row.cells[0].text=='A8':
                before=[c.text for c in row.cells]
                row.cells[2].text='CPU, RSS, Modellgrößen und Queue erfasst; verworfene Entscheidungsfenster gezählt. Kein Nachweis unbegrenzten Dauerbetriebs oder unbekannter physischer Verluste.'
                changes.append({'kind':'requirement_row','before':before,'after':[c.text for c in row.cells],'reason':'Specify software decision-window drops rather than imply all physical losses measured.'});helper.style_table(table)

    anchor=match('Diskussion')
    p=insert_before(anchor,'Tabelle 7-10 fasst die Antworten auf die Forschungsfragen und den Stand der Hypothesen zusammen. Die Aussagen beziehen sich auf die festgelegten Daten und Messgrenzen; aus einer Fensterzahl wird keine zusätzliche Zahl unabhängiger Versuche abgeleitet.')
    p.paragraph_format.keep_with_next=True
    cap=insert_before(anchor,'Tabelle 7-10: Forschungsfragen, Hypothesen und Gültigkeitsgrenzen','Tabellenbeschriftung')
    cap.paragraph_format.keep_with_next=True
    rows=[
        ['Bezug','Beobachteter Befund','Einordnung und Grenze'],
        ['FF1','Anforderungen, Variantenvergleich und realisierte Mess-/Auswertekette dokumentiert.','Auswahl und Implementierung nachvollziehbar; A1 und A11 nur teilweise nachgewiesen. Drehzahl ungemessen.'],
        ['FF2 / H1','F1 der Plattenfolge: RMS 6,93 %, IF 4,04 %, AE 0 %.','H1 in dieser Folge nicht bestätigt. Eine positive Aufnahme; kein allgemeiner Defektnachweis.'],
        ['FF3 / H2','P99 unter dem beobachteten Fensterintervall in allen 18 Sensor-/Replay-Läufen; keine gezählte Überschreitung.','H2 im gemessenen Umfang erfüllt. Latenz ab hostseitig vollständigem Fenster; kein unbegrenzter Echtzeit- oder Dauerbetriebsnachweis.'],
        ['FF4 / H3','50-%-Normaltests mit unverändertem Paket: gepoolte Fehlalarme bei RMS/IF geringer, beim AE geringfügig höher.','H3 nicht pauschal bestätigt. AE-Richtung nur deskriptiv; Startvariabilität und zeitlich getrennte Reihen begrenzen die Zuordnung.']
    ]
    table=doc.add_table(rows=1,cols=3)
    for c,v in zip(table.rows[0].cells,rows[0]):c.text=v
    for values in rows[1:]:
        for c,v in zip(table.add_row().cells,values):c.text=v
    helper.style_table(table);anchor._p.addprevious(table._tbl)
    additions.append({'kind':'table','caption':cap.text,'rows':rows,'anchor':anchor.text})

    insert_after('Die Normaltests zeigen, dass die Normalvalidierung',
        'Die Erkennungskennzahlen sind deskriptive Ergebnisse des vorab definierten Zustandslabels. Für den veränderten Zustand liegt nur eine unabhängige Aufnahme vor; ihre 194 Fenster sind zeitlich abhängig. Ein Konfidenzintervall oder Signifikanztest, der diese Fenster als unabhängige Replikate behandelte, würde die verfügbare Evidenz überschätzen. Deshalb werden Einzelaufnahmen, zeitliche Verläufe und Zählungen berichtet, ohne aus der Fensterzahl eine populationsweite Genauigkeit abzuleiten. Auch der explorative Zusammenhang zwischen Achsenenergie und Modellfehler ist kein kausaler Nachweis.')
    replace('Die lokale Laufzeitprüfung beantwortet eine andere Frage als die Erkennungstests:',
        'Die Laufzeitprüfung zeigt, ob die gespeicherte Verarbeitung im gemessenen Umfang Schritt hält. Numerische Konsistenz und geringe Latenz belegen keine richtige Zustandsentscheidung. Umgekehrt bedeutet schwache Zustandstrennung allein keinen Implementierungsfehler.',
        'Condense repeated distinction between technical execution and recognition, avoiding a nearly empty continuation page.')
    replace('Die Zeitbewertung beginnt mit dem hostseitig vollständigen Fenster.',
        'Die Latenzmessung beginnt beim hostseitig vollständigen Fenster; Fensterfüllzeit und unbekannte Verzögerungen vor dem Host-Leseabschluss sind ausgeschlossen. Endliche Tests unter gewöhnlichem Linux belegen keine harte Echtzeit. Die optionale GUI ist im abschließenden Messweg nicht enthalten; ihre Zusatzlast mit dem neuen Paket wurde nicht geprüft.',
        'Preserve latency endpoints, non-real-time limitation and unmeasured GUI overhead in concise form.')
    replace('Ausblick und verbleibende Nachweise','Nicht erbrachte Nachweise und Ausblick',
        'Present missing measurements as final limitations, not as work claimed completed.')
    replace('Der vorhandene Modellstand sollte als abgeschlossener Pilotstand erhalten bleiben.',
        'Der empirische Teil dieser Fassung wird mit der vorhandenen Datenbasis abgeschlossen. Zwei zusätzliche Plattenfolgen sowie unabhängige Drehzahl- und Sensortaktmessungen wurden vorbereitet, aber nicht durchgeführt. Die damit verbundenen Grenzen bleiben Teil des Ergebnisses: A1 und A11 sind weiterhin nur teilweise nachgewiesen, und die tatsächliche Drehzahl bleibt unbekannt. Die folgende Beschreibung ordnet mögliche Anschlussuntersuchungen ein; sie ist kein Nachweis ihrer Ausführung und verändert den Status des eingefrorenen Entwicklungspakets nicht. Für die wissenschaftliche Dokumentation werden weder weitere Normalstarts zur Verbesserung eines Mittelwerts noch eine nachträgliche Schwellenanpassung vorgenommen.',
        'Explicitly close empirical scope without erasing unmet requirements or pretending future work done.')
    replace('Ein geeigneter Folgeplan würde mehrere vollständige unabhängige',
        'Ein vorbereiteter Anschlussversuch umfasst zwei zusätzliche vollständige Normal–Änderung–Normal-Folgen, sodass zusammen mit der vorhandenen Folge drei unabhängige veränderte Aufnahmen vorlägen. Auch dies wäre ein begrenzter Pilotumfang. Geometrie, Stellvorgabe, Auszeiten, Messdauer, Bewertungsabschnitt und Labels müssten vor Beginn festgelegt bleiben. Die Rückkehrreferenzen würden mit der normalen Variabilität verglichen; genau gleiche RMS-Mittelwerte wären kein notwendiges Rücksetzkriterium. Für weitergehende Aussagen müsste die Zahl unabhängiger Folgen anhand der gewünschten Genauigkeit geplant werden. Weitere Instrumentierung oder Modelländerungen erforderten eine gesondert dokumentierte Versuchs- beziehungsweise Modellversion.',
        'Describe prepared, unperformed follow-up with correct number of independent replicates.')
    replace('FF3 wird durch getrennte Ressourcen- und Latenzmessungen',
        'FF3 wird durch getrennte Ressourcen- und Latenzmessungen im Sensorbetrieb und zeitgetreuen Replay beantwortet. H2 ist im gemessenen Umfang erfüllt: In allen 18 Prozessen liegt P99 unter dem beobachteten Fensterintervall; es gibt keine gezählte Fristüberschreitung. Gemessen wird hostseitige Verarbeitung in 300-s-Prozessen mit 120 s Modellbewertung. Harte Echtzeit, unbegrenzter Dauerbetrieb und fachlich richtige Zustandsentscheidungen folgen daraus nicht.',
        'Explicitly answer H2 with observed latency and its endpoint boundary.')
    replace('FF4 ist durch einen zweiten normalen Betriebspunkt',
        'FF4 ist durch drei unabhängige Normalaufnahmen bei 50 % PWM mit unverändertem Paket untersucht. Gegenüber den festgelegten sechs 75-%-Normalreferenzen sinken die gepoolten Fehlalarmraten von RMS und Isolation Forest; beim Autoencoder steigt der Anteil geringfügig von 5,50 auf 6,01 %. Die Richtung ist nicht über alle Starts einheitlich und belegt keinen isolierten kausalen Tastgradeffekt. H3 wird deshalb nicht pauschal bestätigt.',
        'Keep final H3 answer consistent with run-level evidence.')
    replace('Der wissenschaftliche Beitrag liegt in der konsistenten und überprüfbaren Umsetzung',
        'Der Beitrag liegt in der überprüfbaren Umsetzung und offenen Bewertung eines begrenzten Methodenvergleichs. Stärkere Vibration, höherer Modellscore und richtige Zustandsentscheidung sind verschiedene Sachverhalte. Eine zuverlässige Erkennung der Luftstromänderung ist nicht belegt. Die Pilotstudie ist mit diesem negativen Ergebnis abschließbar; Wiederholungs-, Drehzahl- und physikalische Messkettennachweise bleiben teilweise unerbracht. Vollständige Anforderungserfüllung oder betriebsfertig validierte Defekterkennung werden nicht beansprucht.',
        'Close with concrete contribution, negative result and remaining evidence gaps.')
    replace('Die folgenden Verwechslungszahlen beziehen sich jeweils auf eine vollständige Aufnahme',
        'Die folgenden Verwechslungszahlen sind nach vollständigen Aufnahmen gruppiert, beziehen sich jedoch ausschließlich auf die 194 gültigen Modellfenster im Abschnitt [180,300) s des jeweiligen 300-s-Laufs. Jede Aufnahme enthält eine Zustandsklasse. Der Anteil bezeichnet deshalb entweder Fehlalarme im Normalzustand oder markierte Fenster der kontrollierten Änderung. Kennzahlen der vollständigen Folge stehen in Tabelle A-3.',
        'Clarify full recording as grouping unit versus only prespecified 120-s score interval.')
    replace('Tabelle A-4: Verwechslungszahlen je vollständiger Aufnahme und Methode',
        'Tabelle A-4: Verwechslungszahlen je Aufnahme und Methode im Abschnitt [180,300) s',
        'Avoid suggesting scoring of the entire 300-s acquisition.')

    # Add the digital review provenance without rewriting earlier reports.
    for table in doc.tables:
        if table.rows[0].cells[0].text=='Kennung' and table.rows[0].cells[-1].text=='Projektpfad':
            before=[[c.text for c in row.cells] for row in table.rows]
            table.rows[0].cells[0].text='ID'
            for values in [
                ['B9','Fachreview und unveränderte Zählung je Aufnahme vom 12.09.2026','results/upright_v4_followup_closure_20260912/'],
                ['B10','Abschluss als begrenzte Pilotstudie: Quellen-, Statistik- und Konsistenzprüfung; Änderungs- und Erhaltungsnachweis',str(B.relative_to(ROOT))+'/']
            ]:
                for c,v in zip(table.add_row().cells,values):c.text=v
            helper.style_table(table,is_appendix=True)
            changes.append({'kind':'whole_table','before':before,'after':[[c.text for c in row.cells] for row in table.rows],'reason':'Expose new review provenance as separate digital artifacts.'})
            break
    else:raise AssertionError('Digital-evidence appendix not found')

    appendix_anchor=match('Lesen der Messdateien')
    heading=insert_before(appendix_anchor,'Dokumentation der KI-Unterstützung','Heading 2')
    helper.set_num(heading,num=0)
    insert_before(appendix_anchor,
        'Bei der Bearbeitung wurde Codex von OpenAI als KI-gestütztes Werkzeug für Softwareentwicklung und Fehlersuche, Datenauswertung und Grafikaufbereitung, Quellenprüfung sowie die Formulierung, Überarbeitung und Formatprüfung des Berichts verwendet. Dies umfasst auch generierte und überarbeitete Textpassagen; die Unterstützung beschränkt sich nicht auf Rechtschreibkorrekturen. Die physische Einrichtung und die im Kontrolljournal ausgewiesenen Sichtbeobachtungen stammen aus Nutzerangaben. KI-Ausgaben gelten nicht als eigenständige wissenschaftliche Quellen oder als Ersatz für Messnachweise.')
    insert_before(appendix_anchor,
        'Der digitale Nachweis B10 enthält den Unterstützungsbericht, dokumentierte Änderungen dieser Abschlussfassung sowie die Quellen-, Zahlen- und Erhaltungsprüfungen. Frühere Bearbeitungsartefakte sind unter B7 und B9 zugeordnet. Diese Unterlagen dokumentieren nachvollziehbare Bearbeitungsschritte, jedoch keine lückenlose Passagenzuordnung sämtlicher früherer Fassungen. Eine persönliche fachliche Prüfung, die Eigenständigkeitserklärung und die Einhaltung der für diese Arbeit geltenden Hilfsmittelregeln können nicht durch die Software bestätigt werden.')

    # Additional corrections established by independent source review can be
    # added here as exact, independently logged substitutions.
    corrections=B/'verified_source_corrections.json'
    if corrections.exists():
        for item in json.loads(corrections.read_text()):replace(item['prefix'],item['replacement'],item['reason'])

    doc.core_properties.modified=datetime.now(timezone.utc)
    doc.core_properties.last_modified_by='Codex – Pilotabschluss und Konsistenzprüfung'
    doc.core_properties.revision=(doc.core_properties.revision or 0)+1
    out=B/'thesis_pilot_unformatted.docx';doc.save(out)
    manifest={'utc':datetime.now(timezone.utc).isoformat(),'input_sha256':sha(source),'draft_sha256':sha(out),
              'changes':changes,'additions':additions,'table_count':len(doc.tables),
              'measurement_cutoff':'2026-09-12','new_hardware_measurements':False,'model_or_threshold_changes':False,
              'original_requirements_and_hypotheses_preserved':True,'personal_declaration_and_signature_unchanged':True}
    (B/'word_changes.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps({'draft':str(out),'changes':len(changes),'additions':len(additions),'tables':len(doc.tables)},ensure_ascii=False))

if __name__=='__main__':main()
