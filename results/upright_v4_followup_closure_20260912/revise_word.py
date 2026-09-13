"""Traceable corrections of the preserved final manuscript, no measurement edits."""
from pathlib import Path
from datetime import datetime,timezone
import json,hashlib,importlib.util
from docx import Document
from docx.shared import Pt
B=Path(__file__).resolve().parent;ROOT=B.parents[1];OLD=ROOT/'results/upright_v4_finalization_20260912_103145'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 state=json.loads((B/'review_initial_state.json').read_text());source=Path(state['backup']);assert sha(source)==state['word_sha256']
 from zipfile import ZipFile
 from lxml import etree
 relns='{http://schemas.openxmlformats.org/package/2006/relationships}'
 normalized=B/'normalized_input.docx'
 with ZipFile(source) as zi,ZipFile(normalized,'w') as zo:
  for item in zi.infolist():
   content=zi.read(item)
   if item.filename=='_rels/.rels':
    tree=etree.fromstring(content)
    for rel in tree:
     if rel.get('Type','').endswith('/metadata/core-properties'):
      rel.set('Type','http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties')
    content=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
   zo.writestr(item,content)
 doc=Document(normalized);changes=[]
 def replace_start(prefix,new):
  matches=[p for p in doc.paragraphs if p.text.startswith(prefix)];assert len(matches)==1,(prefix,len(matches))
  p=matches[0];old=p.text;p.text=new;changes.append({'kind':'paragraph','before':old,'after':new})
 replace_start('Der Kernversuch bewertet die lokale Inferenz', 'Der Kernversuch bewertet die lokale Inferenz eines zuvor trainierten Modells auf einer Zielplattform. Die Pipeline bleibt dabei unverändert. Eine mögliche spätere Nachkalibrierung wäre als eigener Versuch auszuwerten. Sie ist nicht Bestandteil des hier abgeschlossenen eingefrorenen Vergleichs.')
 replace_start('Antonini et al. (2023) entwickeln einen ESP32-Sensorknoten', 'Antonini et al. (2023) entwickeln einen ESP32-Sensorknoten für die Überwachung von Unterwasserpumpen. Die Evaluation erfolgt an einer Platine auf einem gefederten PC-Lüfterprüfstand und beschreibt damit keinen gemessenen Pumpenbetrieb. Merkmalsbildung, Isolation-Forest-Training und Inferenz laufen lokal; das Training setzt normalen Betrieb voraus. Die Ressourcenmessungen unterscheiden Laden, Training und Inferenz. Die Erkennungsgenauigkeit wird nicht bewertet.')
 replace_start('H1 setzt einen Vergleich der F1-Werte', 'H1 vergleicht die Erkennungsqualität für die vorab definierten Versuchsanomalien; sie verlangt keine gesicherte Schadensdiagnose. In der vorliegenden Folge ist die äußere Luftstromveränderung das vorab vergebene positive Zustandslabel. Die deskriptiven F1-Werte betragen 6,93 % für RMS, 4,04 % für Isolation Forest und 0 % für den Autoencoder. Die erwartete Überlegenheit des neuronalen Verfahrens ist damit in dieser Folge nicht bestätigt. Das negative Ergebnis ist eine gültige Bewertung des eingefrorenen Pakets. Eine einzige unabhängige veränderte Aufnahme belegt jedoch keine allgemeine Wiederholbarkeit. Die Platte bleibt ein kontrolliert veränderter Betriebszustand und wird nicht als nachgewiesener Defekt bezeichnet.')
 replace_start('Zu FF2 liegen vergleichbare Entscheidungen', 'FF2 wird für den vorab definierten kontrolliert veränderten Betriebszustand durch Kennzahlen auf identischen Testfenstern beantwortet. In der untersuchten Folge beträgt der F1-Wert 6,93 % für RMS, 4,04 % für Isolation Forest und 0 % für den Autoencoder. H1 ist in dieser Folge nicht bestätigt. Die explorative Fehleranalyse erklärt die beobachteten Scores anhand der gespeicherten Achsenskalierung, geringerer X- und Y-Rekonstruktionsfehler im Plattenzustand und Überschneidungen mit bekannter normaler Variation. Die eine veränderte Aufnahme begrenzt die Aussage über Wiederholbarkeit und Übertragbarkeit. Ein allgemeiner Defektnachweis ist nicht Gegenstand dieser abgegrenzten Pilotbewertung; die schlechte Zustandserkennung bleibt als Ergebnis erhalten.')
 replace_start('Die erste offene physikalische Frage betrifft die Zeitbasis.', 'Die erste offene physikalische Frage betrifft die Zeitbasis. Ein kleinster gezielter Versuch wäre eine 60-s-Sensorprobe bei 0 % PWM mit externem Mitschnitt von DATA_READY, SCL und SDA sowie gleichzeitig gespeicherter FIFO-Rohdatei. Gerät, Zeitbasisgenauigkeit und Signalanbindung müssen vorab bekannt sein. Die Interruptausgabe wäre eine eigene dokumentierte Instrumentierungsfassung mit anschließender Wiederherstellung der Register; die bisherige ODR, Auflösung und FIFO-Konfiguration bleiben gleich. DATA_READY kann aktiv bleiben, solange Daten ausstehen. Nur nachgewiesen einzeln aufgelöste Bereitstellungen dürfen zur Periodenmessung dienen; eine Flankenzahl ist nicht automatisch eine Sensorwertzahl. Nicht zuordenbare Abschnitte bleiben erhalten. Aus den vorhandenen Dateien allein ist die genaue Ursache der Abweichung nicht sicher auflösbar.')
 replace_start('Für einen erweiterten Erkennungsnachweis muss zunächst festgelegt werden', 'Die äußere Luftstromveränderung ist als kontrollierte positive Versuchsklasse bereits festgelegt. Für eine Prüfung ihrer Wiederholbarkeit sind weitere vollständige unabhängige Normal–Änderung–Normal-Folgen mit derselben Platte und Geometrie zweckmäßig. Sie würden die vorhandene Folge ergänzen, ohne den Zustand nachträglich umzubenennen oder Schwellen zu ändern. Ein tatsächlicher Defekt muss dafür nicht erzeugt werden. Die Änderung erfolgt ausschließlich bei getrennter externer Versorgung und bestätigtem Stillstand; Sensor und Lüfteraufstellung bleiben unverändert. An einem laufenden Rotor wird nicht gearbeitet.')
 for p in doc.paragraphs:
  needle='Die unabhängigen Normaltests umfassen drei separate Starts bei 75 % PWM.'
  if needle in p.text:
   old=p.text;p.text=p.text.replace(needle,'Die erste unabhängige Normaltestserie umfasst drei separate Starts bei 75 % PWM.');changes.append({'kind':'paragraph','before':old,'after':p.text})
 spec=importlib.util.spec_from_file_location('old_style_helpers',OLD/'build_thesis.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
 cell_changes=0
 for t in doc.tables:
  for row in t.rows:
   if row.cells[0].text=='Antonini et al. (2023)':
    c=row.cells[1];old=c.text;c.text='Unterwasserpumpe als Zielanwendung; Evaluation am gefederten Lüfterprüfstand; Isolation Forest mit Normaltraining';changes.append({'kind':'table_cell','before':old,'after':c.text});helper.style_table(t);cell_changes+=1
   if len(row.cells)==3 and row.cells[0].text=='A9':
    old=[c.text for c in row.cells];row.cells[1].text='Im Pilotumfang ausgewertet';row.cells[2].text='Kennzahlen derselben Zustandsfolge und Verwechslungszahlen je Aufnahme berechnet (Anhang). Geringe Erkennung ist ein gültiges Ergebnis. Nur eine veränderte Aufnahme; Wiederholungsnachweis unter A11 offen.';changes.append({'kind':'requirement_row','before':old,'after':[c.text for c in row.cells]});helper.style_table(t);cell_changes+=1
   if len(row.cells)==3 and row.cells[0].text=='A11':
    old=[c.text for c in row.cells];row.cells[2].text='Aufbau und Protokoll dokumentiert; Normalstarts wiederholt, eine unabhängige Plattenfolge abgeschlossen. Wiederholungen der veränderten Bedingung fehlen. Ein anderer Rückkehrmittelwert allein belegt keine mechanisch misslungene Rücksetzung.';changes.append({'kind':'requirement_row','before':old,'after':[c.text for c in row.cells]});helper.style_table(t);cell_changes+=1
 assert cell_changes==3,cell_changes
 # Add per-record confusion counts before the existing data-file appendix.
 anchor=[p for p in doc.paragraphs if p.text=='Lesen der Messdateien'];assert len(anchor)==1;anchor=anchor[0]
 p=anchor.insert_paragraph_before('Die folgenden Verwechslungszahlen beziehen sich jeweils auf eine vollständige Aufnahme und Methode. Jede Aufnahme enthält eine Zustandsklasse. Der Anteil bezeichnet deshalb entweder Fehlalarme im Normalzustand oder markierte Fenster der kontrollierten Änderung. Kennzahlen der vollständigen Folge stehen in Tabelle A-3.')
 p.paragraph_format.page_break_before=True
 p.paragraph_format.keep_with_next=True
 cap=anchor.insert_paragraph_before('Tabelle A-4: Verwechslungszahlen je vollständiger Aufnahme und Methode',style='Tabellenbeschriftung')
 cap.paragraph_format.keep_with_next=True
 data=json.loads((B/'per_record_metrics.json').read_text())['records'];t=doc.add_table(rows=1,cols=6)
 for c,value in zip(t.rows[0].cells,['Aufnahme / Methode','gültig / ungültig','TP / FN','FP / TN','Anteil [%]','Bedeutung']):c.text=value
 names={'rms':'RMS','isolation_forest':'IF','tflite_autoencoder':'AE'};phases={'normal_before':'Normal vorher','airflow_modified':'Luftstrom verändert','normal_after':'Normal nachher'}
 for r in data:
  values=[phases[r['phase']]+' / '+names[r['method']],f"{r['valid']} / {r['invalid']}",f"{r['tp']} / {r['fn']}",f"{r['fp']} / {r['tn']}",f"{100*r['primary_rate']:.2f}".replace('.',','),'Fehlalarme' if r['label']==0 else 'Änderung markiert']
  for c,value in zip(t.add_row().cells,values):c.text=value
 helper.style_table(t);anchor._p.addprevious(t._tbl)
 anchor.insert_paragraph_before('Die nachträgliche Aufschlüsselung verwendet unveränderte archivierte Entscheidungen. Quellenhashes und Berechnung stehen in results/upright_v4_followup_closure_20260912/per_record_metrics.json; fachliche Korrekturen im dortigen Änderungsnachweis. Die Ausgangsfassung bleibt gesichert.')
 doc.core_properties.modified=datetime.now(timezone.utc);doc.core_properties.last_modified_by='Codex – fachliche Konsistenzprüfung';doc.core_properties.revision=(doc.core_properties.revision or 0)+1
 output=B/'thesis_reviewed_unformatted.docx';doc.save(output)
 manifest={'input_sha256':state['word_sha256'],'output_sha256':sha(output),'changes':changes,'added_table':'Tabelle A-4','existing_frozen_results_unchanged':True,'original_document_backup':str(source),'new_table_count':len(doc.tables),'scope':'Fachliche Konsistenzprüfung; keine persönliche Autoren- oder Betreuerfreigabe'}
 (B/'word_revision_changes.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
 print('Saved',output,'documented existing-text changes:',len(changes),'tables:',len(doc.tables))
if __name__=='__main__':main()
