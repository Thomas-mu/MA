"""Build an evidence-labelled preparation deck from completed technical checks.

This program does not import acquisition code, touch hardware, or use trial data.
Example (python-pptx must be installed in the selected Python environment)::

    python src/build_phone_presentation.py

All generated material and frozen input copies stay below presentation_v1.
The PDF and browser screenshot need LibreOffice and Chromium, respectively.
"""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import html
import json
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "results/phone_vibration_pilot_20260919"
METHODS = (
    ("autoencoder", "Autoencoder", "#007F86", "tflite_autoencoder"),
    ("isolation_forest", "Isolation Forest", "#7454C4", "isolation_forest"),
    ("rms", "RMS", "#D67B19", "rms"),
)
NAVY = "142D42"
MUTED = "587084"
PAPER = "F5F8FA"
TEAL = "007F86"
AMBER = "AB5E06"
WHITE = "FFFFFF"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def fmt(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def snapshot(source: Path, output: Path) -> tuple[dict, dict, list[dict]]:
    """Freeze only the explicitly listed, completed preparation evidence."""
    paths = [source / "plan.json", source / "protocol.md", source / "bundle/bundle.json"]
    records = {}
    for key, _label, _color, _bundle_key in METHODS:
        directory = source / "technical_checks" / key
        report = json.loads((directory / "run.json").read_text())
        summary = json.loads((directory / "summary.json").read_text())
        rows = read_csv(directory / "decisions.csv")
        if report["purpose"] != "technical_check" or report["status"] != "completed":
            raise ValueError(f"Only completed technical checks are allowed: {directory}")
        if summary["call_requests"] != 0 or (directory / "cues.jsonl").stat().st_size:
            raise ValueError("This preparation deck expects no call requests in its sources.")
        if len(rows) != summary["windows"]:
            raise ValueError(f"Window count mismatch: {directory}")
        if len(rows) != 4 or float(report["requested_seconds"]) != 3:
            raise ValueError("The preparation narrative requires the original four-window, 3-s checks.")
        records[key] = {"run": report, "summary": summary, "rows": rows}
        paths.extend(directory / name for name in (
            "run.json", "summary.json", "raw.csv", "decisions.csv", "cues.jsonl",
            "started.json", "raw.events.jsonl", "report.md",
        ))
        if (directory / "capture_source.py").exists():
            paths.append(directory / "capture_source.py")
    sources = []
    for original in paths:
        relative = original.relative_to(source)
        target = output / "sources" / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(original, target)
        digest = sha256(original)
        if sha256(target) != digest:
            raise ValueError(f"Source changed during copy: {original}")
        sources.append({"original": str(original), "snapshot": str(target.relative_to(output)),
                        "sha256": digest, "size_bytes": target.stat().st_size})
    return json.loads((output / "sources/plan.json").read_text()), records, sources


def figure_assets(output: Path, records: dict) -> dict[str, Path]:
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 12,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.labelcolor": f"#{NAVY}", "text.color": f"#{NAVY}",
                         "xtick.color": f"#{MUTED}", "ytick.color": f"#{MUTED}"})
    assets = output / "assets"
    assets.mkdir(exist_ok=True)
    paths = {}

    fig, axes = plt.subplots(1, 3, figsize=(13, 3.15), sharey=True)
    for ax, (key, label, color, _) in zip(axes, METHODS):
        rows = records[key]["rows"]
        ax.plot([int(r["window"]) for r in rows],
                [float(r["score"]) / float(r["threshold"]) for r in rows],
                color=color, marker="o", linewidth=2.5, markersize=7)
        ax.axhline(1, color="#AB5E06", linestyle="--", linewidth=1.5)
        ax.set(title=label, xlabel="Fenster im jeweiligen Einzeltest", xticks=[1, 2, 3, 4],
               ylim=(0, 1.17), yticks=[0, .25, .5, .75, 1])
        ax.grid(alpha=.18, axis="y")
    axes[0].set_ylabel("Score / eigener Schwellenwert")
    fig.tight_layout(pad=1.3)
    for ext in ("png", "pdf"):
        fig.savefig(assets / f"scores.{ext}", dpi=200, facecolor="white")
    plt.close(fig)
    paths["scores"] = assets / "scores.png"

    fig, axes = plt.subplots(1, 2, figsize=(13, 3.15), gridspec_kw={"width_ratios": [1.12, 1]})
    for i, (key, label, color, _) in enumerate(METHODS):
        rows = records[key]["rows"]
        axes[0].plot([int(r["window"]) for r in rows], [float(r["processing_ms"]) for r in rows],
                     color=color, marker="o", linewidth=2, label=label)
        axes[1].scatter([float(r["process_cpu_percent_one_core"]) for r in rows],
                        [i + j * .065 - .10 for j in range(len(rows))],
                        color=color, s=50, alpha=.8)
    axes[0].set(yscale="log", xlabel="Fenster", ylabel="Standardisierung + Entscheidung [ms]",
                xticks=[1, 2, 3, 4], ylim=(.15, 40))
    axes[0].legend(loc="upper right", fontsize=10)
    axes[0].grid(alpha=.2, which="both", axis="y")
    axes[1].set(xlabel="Prozess-CPU [%; 100 % = ein Kern]",
                yticks=[0, 1, 2], yticklabels=[m[1] for m in METHODS], ylim=(2.5, -.5))
    axes[1].grid(alpha=.2, axis="x")
    fig.tight_layout(pad=1.5, w_pad=2.4)
    for ext in ("png", "pdf"):
        fig.savefig(assets / f"resources.{ext}", dpi=200, facecolor="white")
    plt.close(fig)
    paths["resources"] = assets / "resources.png"

    raw = read_csv(output / "sources/technical_checks/autoencoder/raw.csv")
    t = [(int(r["host_monotonic_ns"]) - int(raw[0]["host_monotonic_ns"])) / 1e9 for r in raw]
    fig, ax = plt.subplots(figsize=(12.5, 1.95))
    for field, color in zip(("x_g", "y_g", "z_g"), ("#007F86", "#7454C4", "#D67B19")):
        ax.plot(t, [float(r[field]) for r in raw], lw=.9, color=color, label=field[0])
    ax.set(xlabel="Host-Lesezeit relativ zum ersten Sample [s]", ylabel="Beschleunigung [g]")
    ax.legend(ncol=3, loc="upper right", fontsize=10)
    ax.grid(alpha=.15)
    fig.tight_layout(pad=.7)
    for ext in ("png", "pdf"):
        fig.savefig(assets / f"raw_example.{ext}", dpi=200, facecolor="white")
    plt.close(fig)
    paths["raw"] = assets / "raw_example.png"
    return paths


def make_report(output: Path, records: dict, timestamp: str) -> Path:
    rows = []
    for key, label, color, _ in METHODS:
        summary = records[key]["summary"]
        rows.append(f'<tr><td><span style="color:{color}">●</span> {html.escape(label)}</td>'
                    f'<td>{summary["windows"]}</td><td>{summary["invalid_windows"]}</td>'
                    f'<td>{summary["alarm_windows"]}</td><td>0</td></tr>')
    markup = '''<!doctype html><html lang="de"><meta charset="utf-8">
<title>Vorbereitungsstand | Handyvibration auf dem Raspberry Pi</title>
<style>*{box-sizing:border-box}body{margin:0;padding:36px 54px;background:#F5F8FA;color:#142D42;
font-family:"DejaVu Sans",sans-serif}header{border-bottom:3px solid #007F86;padding-bottom:20px}
.eyebrow{color:#007F86;letter-spacing:2px;font-size:15px;font-weight:700}h1{font-size:34px;margin:12px 0}
.sub{font-size:18px;color:#587084}.grid{display:grid;grid-template-columns:1fr 1fr;gap:24px;margin-top:24px}
.card{background:white;border:1px solid #DCE6EC;border-radius:12px;padding:24px}h2{font-size:23px;margin:0 0 14px}
table{border-collapse:collapse;width:100%;font-size:17px}th,td{text-align:left;padding:13px 9px;border-bottom:1px solid #DCE6EC}
th{font-size:13px;color:#587084}p{font-size:17px;line-height:1.55;margin:12px 0}
.tag{display:inline-block;background:#FFF1DB;color:#87500F;padding:8px 12px;border-radius:6px;font-size:14px;font-weight:bold}
img{width:100%;display:block}.note{font-size:14px;color:#587084}.full{margin-top:24px}
footer{margin-top:20px;font-size:13px;color:#587084}</style>
<header><div class="eyebrow">EDGE AI / DOKUMENTIERTER VORBEREITUNGSSTAND</div>
<h1>Handyvibration: drei Methoden, ein gemeinsames Prüfkonzept</h1>
<div class="sub">Raspberry Pi 5 · 75 % PWM · 128 × 3 Werte je Fenster · 200 Hz nominal</div></header>
<div class="grid"><section class="card"><h2>Abgeschlossene Techniktests</h2>
<table><tr><th>METHODE</th><th>FENSTER</th><th>UNGÜLTIG</th><th>ALARME</th><th>ANRUFE*</th></tr>'''
    markup += "".join(rows)
    markup += '''</table><p class="note">Jeweils 3 s angefordert; unabhängige kurze Aufnahmen.
* Keine Anrufaufforderung protokolliert. Keine Störungsversuche in dieser Datengrundlage.</p></section>
<section class="card"><h2>Was diese Daten belegen</h2><p>Die drei Einzelprozesse können Sensordaten
aufzeichnen, Scores berechnen und Entscheidungen speichern.</p>
<span class="tag">Keine Rangfolge der Fehlererkennung</span><p>Physischer Vibrationsbeginn,
endgültiger Handyaufbau und dessen Kalibrierung sind noch offen.</p></section></div>
<section class="card full"><h2>Scoreverlauf der vier Fenster je Techniktest</h2>
<img src="assets/scores.png" alt="Drei Scoreverläufe relativ zu ihrer eigenen Schwelle">
<p class="note">Gestrichelt: eigener Schwellenwert. Score/Schwelle ist keine Wahrscheinlichkeit
und macht unterschiedliche Scores nicht zu gleicher Empfindlichkeit.</p></section>
<footer>Nachträglich erzeugter Evidenzbericht · keine Live-GUI · keine Aufnahme des Versuchsaufbaus.<br>
Quellen: sources/technical_checks/*/{run.json, decisions.csv, summary.json} · Erstellt: '''
    markup += html.escape(timestamp) + "</footer></html>"
    path = output / "evidence_report.html"
    path.write_text(markup)
    return path


def screenshot_report(report: Path, output: Path) -> Path:
    screenshot = output / "assets/evidence_report_screenshot.png"
    with tempfile.TemporaryDirectory(prefix="edgeai-report-browser-") as profile:
        command = ["chromium", "--headless", "--no-sandbox", "--disable-gpu",
                   "--disable-dev-shm-usage", "--no-first-run", "--hide-scrollbars",
                   f"--user-data-dir={profile}", "--window-size=1600,1100",
                   "--force-device-scale-factor=1", "--virtual-time-budget=1500",
                   f"--screenshot={screenshot}", report.as_uri()]
        completed = subprocess.run(command, text=True, capture_output=True, timeout=90)
    (output / "browser_capture.log").write_text(completed.stdout + completed.stderr)
    completed.check_returncode()
    with Image.open(screenshot) as captured:
        if captured.size != (1600, 1100):
            raise ValueError(f"Unexpected screenshot dimensions: {captured.size}")
    (output / "browser_capture.json").write_text(json.dumps({
        "kind": "real_headless_browser_screenshot_of_local_evidence_report",
        "report": report.name, "report_sha256": sha256(report),
        "screenshot": str(screenshot.relative_to(output)), "sha256": sha256(screenshot),
        "viewport": [1600, 1100], "desktop_captured": False,
        "live_gui_captured": False, "physical_setup_photographed": False,
        "captured_utc": datetime.now(timezone.utc).isoformat(),
    }, ensure_ascii=False, indent=2) + "\n")
    return screenshot


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333333), Inches(7.5)
        self.notes = []
        self.prs.core_properties.title = "Handyvibration auf dem Raspberry Pi – Vorbereitungsstand"
        self.prs.core_properties.subject = "Techniktest, Versuchsdesign und Evidenzgrenzen"
        self.prs.core_properties.author = "Masterarbeit Edge AI"

    def text(self, slide, x, y, w, h, text, size=20, color=NAVY, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.015)
        frame.margin_top = frame.margin_bottom = Inches(.01)
        for i, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "DejaVu Sans"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(8)
        return box

    def box(self, slide, x, y, w, h, fill=WHITE, line=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background() if line is None else None
        if line:
            shape.line.color.rgb = RGBColor.from_string(line)
        shape.adjustments[0] = .10
        return shape

    def picture(self, slide, path, x, y, w, h):
        with Image.open(path) as image:
            ratio = image.width / image.height
        actual_w, actual_h = (w, w / ratio) if w / h <= ratio else (h * ratio, h)
        return slide.shapes.add_picture(str(path), Inches(x + (w - actual_w) / 2),
                                       Inches(y + (h - actual_h) / 2),
                                       width=Inches(actual_w), height=Inches(actual_h))

    def slide(self, eyebrow, title, subtitle, notes):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        self.text(slide, .55, .30, 11.9, .3, eyebrow.upper(), 11, TEAL, True)
        self.text(slide, .55, .81, 12.2, .7, title, 30, NAVY, True)
        self.text(slide, .55, 1.58, 12.1, .66, subtitle, 15, MUTED)
        index = len(self.prs.slides)
        self.text(slide, .55, 7.09, 11.7, .24,
                  "VORBEREITUNGSSTAND · 19.09.2026 · Techniktests, keine Ergebnisse von Anrufversuchen", 9, MUTED)
        self.text(slide, 12.22, 7.04, .6, .30, f"{index:02d}", 12, TEAL, True)
        slide.notes_slide.notes_text_frame.text = notes
        self.notes.append((title, notes))
        return slide


def build_deck(output: Path, plan: dict, records: dict, assets: dict, screenshot: Path) -> Path:
    d = Deck()
    s = d.slide("Forschungsfrage", "Wie früh erkennen drei Methoden eine Vibration?",
                "Versuchsplanung und technische Vorbereitung für einen Vergleich auf dem Raspberry Pi 5.",
                "Diese Fassung verwendet ausschließlich die abgeschlossenen technischen Kurztests vom 19.09.2026. "
                "Reale Anrufversuche wurden inzwischen begonnen; diese separate Vorbereitungsfassung wertet sie nicht aus. "
                "Sie enthält keine Resultate eines Handy-Anrufversuchs. Eine externe Handyvibration ist eine gezielt "
                "eingebrachte Störung, kein Nachweis eines echten Lüfterdefekts. Quellen: plan.json; protocol.md; "
                "technical_checks/*/run.json. Der gewählte Betriebspunkt beträgt 75 % PWM, nicht gemessene Drehzahl.")
    d.box(s, .55, 2.5, 7.25, 3.9)
    d.text(s, .9, 2.85, 6.5, 1.08, "Erkennung und Ressourcen\ngemeinsam bewerten", 27, NAVY, True)
    d.text(s, .9, 4.15, 6.25, 1.62, "Autoencoder · Isolation Forest · RMS\nEinzelbetrieb je Methode\nRohdaten, Entscheidungen und Systemwerte archivieren", 21)
    for i, (big, small) in enumerate((("3", "Methoden"), ("75 %", "PWM-Betriebspunkt"), ("0", "Anrufversuche in diesen Quellen"))):
        y = 2.5 + i * 1.32
        d.box(s, 8.1, y, 4.68, 1.13, "E6F2F1")
        d.text(s, 8.4, y + .12, 1.5, .65, big, 29, TEAL, True)
        d.text(s, 10.0, y + .16, 2.5, .76, small, 15)

    s = d.slide("Versuchsdesign", "Ein Anruf pro Methode: zunächst ein Pilotversuch",
                "Gleicher Aufbau, gleiche Einstellungen; die drei Live-Anrufe bleiben physisch unterschiedliche Ereignisse.",
                "Plan: ein Anruf je Methode, insgesamt drei Anrufe. Vorlauf mindestens 30 s, Zielvibration etwa 5 s, "
                "anschließend mindestens 15 s Ruhe. Anrufsignal und Benutzerbestätigung werden separat protokolliert. "
                "Die tatsächliche Vibration kann später beginnen und anders lang dauern. Telefonposition, Befestigung, "
                "Vibrationsmuster und Klingelton aus sind vorab zu dokumentieren. Ein Ereignis pro Methode erlaubt "
                "keine belastbare allgemeine Erkennungsrate. Für exakt identische Eingaben ist Replay notwendig.")
    stages = [("01", "Normalphase", "≥ 30 s\nHandy liegt ruhig"),
              ("02", "Bestätigung", "Auf Aufforderung\nAnruf ankündigen"),
              ("03", "Vibration", "Ziel: etwa 5 s\nein Anruf"),
              ("04", "Nachlauf", "≥ 15 s\nEnde protokollieren")]
    for i, (n, title, detail) in enumerate(stages):
        x = .55 + i * 3.1
        d.box(s, x, 2.65, 2.90, 2.23)
        d.text(s, x + .20, 2.88, 2.5, .35, n, 14, TEAL, True)
        d.text(s, x + .20, 3.42, 2.5, .48, title, 20, NAVY, True)
        d.text(s, x + .20, 4.06, 2.5, .7, detail, 16)
    d.box(s, .55, 5.3, 12.2, 1.2, "FFF1DB")
    d.text(s, .83, 5.50, 11.65, .82,
           "Fairer Methodenvergleich: gespeicherte Aufnahme später allen Methoden identisch vorspielen.\nDie Live-Einzelversuche prüfen zusätzlich den tatsächlichen Betrieb auf dem Pi.", 18)

    s = d.slide("Datenfluss", "Gemeinsame Messkette, eine aktive Methode",
                "ADXL345 · nominal 200 Hz · 128 × 3 Werte pro Fenster · nominale Fensterperiode 0,64 s.",
                "Die Pipeline verwendet den ADXL345 im FIFO-Betrieb. Fenstergröße und Schrittweite betragen 128. "
                "Die nominale Periode ist 128/200 = 0,64 s; Host-Lesezeitstempel sind keine direkt gemessenen "
                "Sensor-Abtastzeitpunkte. Standardisierung erfolgt mit den eingefrorenen Trainingsparametern. "
                "Im Live-Einzelbetrieb wird eine Methode ausgeführt. Das Beispiel unten zeigt echte Rohdaten "
                "des abgeschlossenen Autoencoder-Techniktests, ohne Anruf. Rohdatenquelle: "
                "sources/technical_checks/autoencoder/raw.csv. Keine Störung in diesem Bild behauptet.")
    for i, (title, detail) in enumerate((("Sensor", "x, y, z in g"), ("Fenster", "128 × 3 Werte"),
                                       ("Skalierung", "Trainingsparameter"), ("Eine Methode", "Score > Schwelle"),
                                       ("Protokoll", "CSV + JSON"))):
        x = .55 + i * 2.47
        d.box(s, x, 2.47, 2.25, 1.03, "E6F2F1" if i == 3 else WHITE)
        d.text(s, x + .13, 2.62, 1.98, .38, title, 18, NAVY, True)
        d.text(s, x + .13, 3.07, 1.98, .30, detail, 12, MUTED)
        if i < 4:
            d.text(s, x + 2.28, 2.83, .25, .38, "→", 17, TEAL)
    d.text(s, .62, 3.88, 12, .35, "Echte Beispielaufnahme aus dem Autoencoder-Techniktest", 16, NAVY, True)
    d.picture(s, assets["raw"], .55, 4.28, 12.2, 2.3)

    s = d.slide("Entscheidungsregel", "Eigene Scores, eigene Schwellenwerte",
                "In allen drei Fällen gilt: Score > Schwelle → Alarm. Die Scorewerte sind nicht direkt untereinander vergleichbar.",
                "Die Standardisierung erfolgt achsweise. Autoencoder: mittlere quadratische Abweichung zwischen "
                "Eingabe und Rekonstruktion über 384 Werte. Isolation Forest: negatives score_samples auf dem "
                "flachgelegten standardisierten Fenster. RMS: Wurzel des mittleren Quadrats des standardisierten "
                "Fensters. P99 bedeutet 99. Perzentil der normalen Validierungsscores, nicht 99 % Genauigkeit. "
                "Der AE-Schwellwert ist aus der eingefrorenen Keras-Kalibrierung übernommen, die anderen aus dem "
                "Vergleichsbündel. Es gibt 97 Validierungsfenster; die AE-Validierung wurde zusätzlich zur "
                "Trainingsüberwachung/Early Stopping verwendet, kein eigener dritter Kalibrierungssplit. "
                "Quelle: bundle/bundle.json, plan.json und archivierte Vergleichsimplementierung.")
    descriptions = ("Rekonstruktionsfehler\nMSE über 384 Werte", "Ausreißerbewertung\n−score_samples", "Signalstärke\nRMS des skalierten Fensters")
    for i, ((key, label, color, bundle_key), detail) in enumerate(zip(METHODS, descriptions)):
        x = .55 + i * 4.15
        d.box(s, x, 2.65, 3.91, 2.3)
        d.text(s, x + .22, 2.91, 3.47, .45, label, 21, color[1:], True)
        d.text(s, x + .22, 3.56, 3.48, .83, detail, 17)
        d.text(s, x + .22, 4.45, 3.47, .34, "Schwelle " + fmt(plan["thresholds"][bundle_key]["value"], 6), 17, NAVY, True)
    d.text(s, .62, 5.38, 12, .55, "P99 aus normaler Validierung: ein Kalibrierungsprinzip, keine Gütegarantie.", 21, TEAL, True)
    d.text(s, .62, 6.07, 11.9, .56, "97 Validierungsfenster · AE-Validierung auch für Trainingsüberwachung genutzt · neuer Handyaufbau noch ungeprüft", 15, MUTED)

    s = d.slide("Echte Techniktestdaten", "Zwölf Fenster verarbeitet, kein Alarm",
                "Drei getrennte Kurztests, jeweils 3 s angefordert und vier vollständige Fenster. Keine protokollierten Anrufereignisse.",
                "Jedes Teilbild zeigt einen anderen technischen Test. Für bessere Lesbarkeit sind Scores durch "
                "die jeweils eigene positive Schwelle geteilt. Der Wert 1 markiert diese Schwelle. Die Division "
                "liefert keine kalibrierte Wahrscheinlichkeit und erlaubt keine Aussage über gleiche Sensitivität. "
                "Alle 12 vollständigen Fenster sind gültig und als normal klassifiziert. Das belegt die Funktionsfähigkeit "
                "der Pipeline in diesen kurzen Aufnahmen, keine getestete Fehlererkennung oder Fehlalarmrate. "
                "Quellen: technical_checks/*/decisions.csv und summary.json.")
    d.picture(s, assets["scores"], .55, 2.65, 12.2, 3.17)
    d.box(s, .55, 6.07, 12.2, .67, "E6F2F1")
    d.text(s, .78, 6.22, 11.75, .38, "Befund: Aufnahme und Entscheidung funktionieren. Eine Rangfolge der Erkennung ist daraus nicht ableitbar.", 16, NAVY, True)

    s = d.slide("Systemmessung", "Berechnung und Prozesslast im Techniktest",
                "Vier Beobachtungen je Methode dienen hier nur dem technischen Nachweis; dies ist kein belastbarer Ressourcenbenchmark.",
                "processing_ms misst laut archiviertem Erfassungscode Standardisierung plus Scoreberechnung und "
                "Klassifikation nach Entnahme des Fensters. Es ist weder reine Modell-Inferenz noch physische "
                "Fehlererkennungsverzögerung. CPU und RAM gelten für den gesamten Einzelprozess einschließlich "
                "Sensorerfassung und Logging. CPU 100 % entspricht einem Kern. RAM sind abgetastete RSS-Werte "
                "während des Laufs, nicht der Gesamtstart-Peak. Die technischen Tests wurden nacheinander und teils "
                "mit unterschiedlichen Codeversionen durchgeführt; vier Fenster sind kein seriöser Benchmark. "
                "Keine Rangfolge der Methodenqualität aus dieser Folie ziehen. Quellen: decisions.csv; run.json.")
    d.picture(s, assets["resources"], .55, 2.50, 12.2, 3.04)
    for i, (key, label, color, _) in enumerate(METHODS):
        x = .63 + i * 4.13
        summary = records[key]["summary"]
        d.text(s, x, 5.87, 3.9, .30, label, 16, color[1:], True)
        d.text(s, x, 6.26, 3.90, .52,
               f"Median: {fmt(summary['processing_median_ms'], 3)} ms\nMax. abgetastete RSS: {fmt(summary['sampled_peak_rss_mib'], 1)} MiB", 13)

    s = d.slide("Zeitmessung richtig lesen", "Anrufsignal und Vibrationsbeginn unterscheiden",
                "Für eine echte Erkennungsverzögerung brauchen wir einen unabhängigen Zeitbezug zur physischen Störung.",
                "Cue-Zeitstempel dokumentieren die Interaktion, nicht den Vibrationsmotor des Telefons. Ein "
                "erster Alarm relativ zu einer Aufforderung wäre nur ein deskriptiver Cue-Abstand. Physischer Beginn "
                "und Ende müssen separat markiert oder gemessen werden; eine nachträgliche Schätzung aus demselben "
                "Signal ist als Schätzung zu kennzeichnen. Beim Replay derselben Aufnahme kann man bereits das "
                "erste Alarmfenster vergleichen, ohne exakte echte Latenz zu behaupten. Gleiches Fenster bedeutet "
                "Gleichstand auf Fensterauflösung. Anrufverzögerung und Fenstertakt dürfen nicht als Modellzeit "
                "fehlinterpretiert werden.")
    items = [("Aufforderung", "Chat / Software-Cue", TEAL), ("Vibrationsbeginn", "Physisch noch ungemessen", AMBER),
             ("Fenster vollständig", "Messdaten liegen vor", TEAL), ("Erster Alarm", "Entscheidung protokolliert", TEAL)]
    for i, (title, detail, color) in enumerate(items):
        x = .55 + i * 3.1
        d.box(s, x, 2.70, 2.9, 1.72, "FFF1DB" if i == 1 else WHITE)
        d.text(s, x + .16, 2.95, 2.6, .58, title, 19, color, True)
        d.text(s, x + .16, 3.70, 2.58, .50, detail, 14)
    d.text(s, .68, 4.94, 12.0, .70, "Erkennungsverzögerung = erster Alarm − tatsächlicher Vibrationsbeginn", 23, NAVY, True)
    d.text(s, .68, 5.98, 11.9, .68, "Schon vergleichbar im Replay: erstes Alarmfenster auf identischen Daten.\nNoch offen: belastbare Verzögerung ab der tatsächlichen Störung.", 18, MUTED)

    s = d.slide("Nachvollziehbare Dokumentation", "Echter Screenshot eines lokalen Evidenzberichts",
                "Der Browser zeigt die archivierten Techniktestdaten. Das Bild ist keine Live-GUI und kein Foto des Aufbaus.",
                "Der abgebildete Screenshot wurde tatsächlich mit Chromium aus evidence_report.html aufgenommen. "
                "Er zeigt nur den hier erzeugten lokalen Bericht; weder Desktop noch private Anwendungen wurden "
                "aufgenommen. Rohdaten und Metadaten liegen als eingefrorene Kopien im sources-Verzeichnis. "
                "source_manifest.json enthält SHA-256-Hashes; browser_capture.json dokumentiert Zeitpunkt, Quelle "
                "und Screenshot-Hash. Die übrigen Diagramme sind aus CSV erzeugte Exporte und werden nicht als "
                "Screenshots bezeichnet. Ein Foto der physischen Handyposition ist weiterhin nicht vorhanden.")
    d.picture(s, screenshot, .55, 2.39, 7.62, 4.46)
    d.box(s, 8.48, 2.63, 4.30, 3.73)
    d.text(s, 8.74, 2.95, 3.77, .45, "Prüfbare Artefakte", 22, NAVY, True)
    d.text(s, 8.74, 3.65, 3.72, 2.36,
           "Rohdaten und Entscheidungen\nLauf- und Modellmetadaten\nPNG-/PDF-Diagramme\nBrowser-Screenshot\nSHA-256-Quellenmanifest", 17)

    s = d.slide("Aussagekraft", "Technisch vorbereitet; Nachweise bleiben offen",
                "Die Folien trennen vorhandene Evidenz von dem, was die Anrufversuche erst liefern sollen.",
                "Die Kalibrierungsdaten stammen aus dem bisherigen Standlaufprofil bei 75 % PWM. Ein angebrachtes "
                "Handy verändert möglicherweise die mechanische Ankopplung; der endgültige Aufbau ist noch nicht "
                "bestätigt. Das Bundle trägt measurement_chain_status=unverified_legacy, weshalb hier kein "
                "abschließend validiertes Messsystem behauptet wird. Normaldaten für Aufbaukontrolle und ggf. neue "
                "Kalibrierung sind vor den bewerteten Tests notwendig. Ein einziger Anruf pro Methode ist ein Pilot, "
                "kein Zuverlässigkeitsnachweis. Externe Tischvibrationen sind keine echten Lüfterdefekte.")
    rows = [("Einzelbetrieb, Rohdaten, Scores", "in drei kurzen Tests belegt", TEAL),
            ("Handyposition und finale Normalkalibrierung", "noch zu bestätigen / prüfen", AMBER),
            ("Physischer Störungsbeginn und Störungsende", "noch nicht unabhängig gemessen", AMBER),
            ("Erkennungsrate, Fehlalarme, Ressourcenrangfolge", "noch keine belastbare Aussage", AMBER)]
    for i, (topic, status, color) in enumerate(rows):
        y = 2.58 + i * .86
        d.box(s, .55, y, 12.2, .69)
        d.text(s, .78, y + .16, 7.0, .38, topic, 17, NAVY, True)
        d.text(s, 8.00, y + .16, 4.42, .38, status, 15, color)
    d.text(s, .65, 6.30, 12.0, .37, "Untersuchungsgegenstand: extern eingebrachte Vibration — kein belegter Lüfterdefekt.", 16, MUTED)

    s = d.slide("Nächste Schritte", "Vom Probelauf zur auswertbaren Versuchsreihe",
                "Aufnahme und Auswertung erfolgen ohne Live-GUI; die Bedienperson löst auf Signal genau einen Anruf aus.",
                "Als nächstes Aufbau dokumentieren und Normalzustand prüfen. Bei verändertem mechanischem Aufbau "
                "neue Normaldaten aufnehmen und Kalibrierung erneuern. Dann jeden Einzelprozess separat aufnehmen, "
                "Aufforderung, Anrufbestätigung und Ende protokollieren, Rohdaten und Fehler erhalten. Tatsächliche "
                "physische Zeitreferenz ist für Latenz nötig. Danach gleiche Aufnahme für alle Methoden auswerten. "
                "Ressourcenbenchmark separat mit fester Umgebung, Aufwärmphase, ausreichend langer Laufzeit und "
                "Wiederholungen durchführen. Ergebnisse der Pilotversuche werden in einer neuen Präsentationsversion "
                "ergänzt; diese Version bleibt als Vorbereitungsevidenz unverändert.")
    next_steps = [("01", "Aufbau festhalten", "Handy befestigen, Vibration einstellen, Normalzustand prüfen."),
                  ("02", "Einzelversuche aufnehmen", "Je Methode ein Anruf, Bestätigung und Ende protokollieren."),
                  ("03", "Identische Daten auswerten", "Replay: Alarmfenster vergleichen; Zeitreferenz transparent behandeln."),
                  ("04", "Ergebnisse präsentieren", "Diagramme, Einzelereignisse und Grenzen; später Wiederholungen.")]
    for i, (number, title, detail) in enumerate(next_steps):
        y = 2.47 + i * .97
        d.text(s, .66, y + .07, .65, .43, number, 23, TEAL, True)
        d.text(s, 1.60, y, 10.77, .43, title, 21, NAVY, True)
        d.text(s, 1.60, y + .46, 10.77, .43, detail, 16, MUTED)

    path = output / "handyvibration_vorbereitung.pptx"
    d.prs.save(path)
    explanation = "# Erläuterungen zur Präsentation\n\n"
    explanation += "Stand: Vorbereitung; ausschließlich abgeschlossene Techniktests, keine Anrufversuchsergebnisse.\n\n"
    for i, (title, notes) in enumerate(d.notes, 1):
        explanation += f"## Folie {i}: {title}\n\n{notes}\n\n"
    (output / "speaker_notes.md").write_text(explanation)
    return path


def html_slide_export(pptx: Path, output: Path) -> Path:
    """Render this deck's editable shapes/images when Impress is unavailable."""
    presentation = Presentation(pptx)
    emu_to_px = 96 / 914400
    width, height = presentation.slide_width * emu_to_px, presentation.slide_height * emu_to_px
    parts = [f'''<!doctype html><html lang="de"><meta charset="utf-8"><title>{html.escape(presentation.core_properties.title or 'Präsentation')}</title>
<style>@page{{size:{width}px {height}px;margin:0}}*{{box-sizing:border-box}}
html,body{{padding:0;margin:0}}body{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.slide{{position:relative;width:{width}px;height:{height}px;overflow:hidden;page-break-after:always}}
.slide:last-child{{page-break-after:auto}}p{{margin:0;line-height:1.08}}img{{display:block}}</style>''']
    for slide in presentation.slides:
        parts.append(f'<section class="slide" style="background:#{slide.background.fill.fore_color.rgb}">')
        for shape in slide.shapes:
            style = f"position:absolute;left:{shape.left * emu_to_px}px;top:{shape.top * emu_to_px}px;"
            style += f"width:{shape.width * emu_to_px}px;height:{shape.height * emu_to_px}px;"
            if shape.shape_type == 13:  # PICTURE
                encoded = base64.b64encode(shape.image.blob).decode()
                parts.append(f'<img style="{style}" src="data:{shape.image.content_type};base64,{encoded}">')
            elif shape.has_text_frame and any(p.text for p in shape.text_frame.paragraphs):
                frame = shape.text_frame
                style += f"padding:{frame.margin_top * emu_to_px}px {frame.margin_right * emu_to_px}px "
                style += f"{frame.margin_bottom * emu_to_px}px {frame.margin_left * emu_to_px}px;"
                parts.append(f'<div style="{style}">')
                for paragraph in frame.paragraphs:
                    font = paragraph.font
                    size = font.size.pt if font.size else 20
                    color = str(font.color.rgb) if font.color.type else NAVY
                    after = paragraph.space_after.pt if paragraph.space_after else 0
                    parts.append(f'<p style="font-family:DejaVu Sans,sans-serif;font-size:{size}pt;'
                                 f'font-weight:{700 if font.bold else 400};color:#{color};margin-bottom:{after}pt">'
                                 f'{html.escape(paragraph.text)}</p>')
                parts.append('</div>')
            elif getattr(shape, "fill", None) is not None and shape.fill.type:
                parts.append(f'<div style="{style}background:#{shape.fill.fore_color.rgb};border-radius:9px"></div>')
        parts.append('</section>')
    parts.append('</html>')
    page = output / "presentation_print.html"
    page.write_text("\n".join(parts))
    return page


def export_pdf_and_previews(pptx: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="edgeai-lo-profile-") as profile:
        completed = subprocess.run([
            "libreoffice", "--headless", f"-env:UserInstallation={Path(profile).as_uri()}",
            "--convert-to", "pdf", "--outdir", str(output), str(pptx),
        ], capture_output=True, text=True, timeout=120)
    (output / "pdf_export.log").write_text(completed.stdout + completed.stderr)
    pdf = pptx.with_suffix(".pdf")
    if (completed.returncode != 0 or "source file could not be loaded" in completed.stdout + completed.stderr
            or not pdf.is_file() or pdf.stat().st_mtime < pptx.stat().st_mtime):
        page = html_slide_export(pptx, output)
        with tempfile.TemporaryDirectory(prefix="edgeai-slides-browser-") as profile:
            fallback = subprocess.run([
                "chromium", "--headless", "--no-sandbox", "--disable-gpu", "--no-first-run",
                f"--user-data-dir={profile}", "--no-pdf-header-footer", "--virtual-time-budget=1500",
                f"--print-to-pdf={pdf}", page.as_uri(),
            ], text=True, capture_output=True, timeout=90)
        (output / "pdf_export.log").write_text(completed.stdout + completed.stderr +
                                               "\nChromium fallback:\n" + fallback.stdout + fallback.stderr)
        fallback.check_returncode()
        engine = "Chromium: HTML print view reconstructed from PPTX text, geometry and images (Impress unavailable)"
    else:
        engine = "LibreOffice Impress"
    if not pdf.is_file():
        raise RuntimeError("No PDF was created")
    (output / "pdf_export.json").write_text(json.dumps({"engine": engine,
        "note": "PDF is a separately rendered export; text may wrap differently in PowerPoint."}, indent=2) + "\n")
    previews = output / "previews"
    previews.mkdir(exist_ok=True)
    subprocess.run(["pdftoppm", "-scale-to", "1600", "-png", str(pdf), str(previews / "slide")],
                   check=True, capture_output=True, timeout=120)
    images = sorted(previews.glob("slide-*.png"))
    expected_count = len(Presentation(pptx).slides)
    if len(images) != expected_count:
        raise ValueError(f"Expected {expected_count} slide previews; found {len(images)}")
    contact = Image.new("RGB", (1600, ((expected_count + 1) // 2) * 450), "#DCE6EC")
    for i, image_path in enumerate(images):
        with Image.open(image_path) as rendered:
            rendered.thumbnail((790, 440))
            contact.paste(rendered, ((i % 2) * 800 + 5, (i // 2) * 450 + 5))
    contact.save(previews / "contact_sheet.jpg", quality=92)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_INPUT)
    args = parser.parse_args()
    source = args.source.resolve()
    output = source / "presentation_v1"
    output.mkdir(exist_ok=True)
    created = datetime.now(timezone.utc).isoformat()
    plan, records, sources = snapshot(source, output)
    generator_snapshot = output / "sources/build_phone_presentation.py"
    shutil.copy2(Path(__file__), generator_snapshot)
    sources.append({"original": str(Path(__file__)),
                    "snapshot": str(generator_snapshot.relative_to(output)),
                    "sha256": sha256(generator_snapshot),
                    "size_bytes": generator_snapshot.stat().st_size})
    (output / "source_manifest.json").write_text(json.dumps({
        "created_utc": created, "scope": "completed_technical_checks_only_no_phone_trials",
        "script_sha256": sha256(Path(__file__)), "sources": sources,
    }, ensure_ascii=False, indent=2) + "\n")
    assets = figure_assets(output, records)
    report = make_report(output, records, created)
    screenshot = screenshot_report(report, output)
    pptx = build_deck(output, plan, records, assets, screenshot)
    export_pdf_and_previews(pptx, output)
    (output / "README.md").write_text(
        "# Präsentation: Vorbereitungsstand\n\n"
        "Diese Fassung enthält ausschließlich drei abgeschlossene Techniktests (je vier Fenster), "
        "keine Ergebnisse der Handy-Anrufversuche. Reale Anrufversuche wurden inzwischen begonnen; "
        "sie werden separat ausgewertet und sind nicht Teil dieser Datenbasis.\n\n"
        "- handyvibration_vorbereitung.pptx: editierbare 16:9-Präsentation mit Sprechernotizen.\n"
        "- handyvibration_vorbereitung.pdf: exportierte Präsentation.\n"
        "- speaker_notes.md: Erläuterungen, Definitionen und Quellen je Folie.\n"
        "- evidence_report.html: lokaler Evidenzbericht.\n"
        "- assets/evidence_report_screenshot.png: echter Browser-Screenshot dieses Berichts; keine Live-GUI.\n"
        "- assets/: Diagramme als PNG und PDF.\n"
        "- sources/: eingefrorene Quellenkopien, ohne laufende Versuche.\n"
        "- source_manifest.json / artifact_manifest.json: SHA-256-Nachweise.\n"
        "- previews/: gerenderte Folien zur Sichtprüfung.\n\n"
        "Die Aufnahmesoftware und Hardware werden vom Generator nicht aufgerufen. "
        "Rendern benötigt CPU/RAM und soll außerhalb bewerteter Ressourcenmessungen stattfinden.\n"
    )
    artifact_files = sorted(p for p in output.rglob("*") if p.is_file() and p.name != "artifact_manifest.json")
    (output / "artifact_manifest.json").write_text(json.dumps({
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "files": [{"path": str(p.relative_to(output)), "sha256": sha256(p), "size_bytes": p.stat().st_size}
                  for p in artifact_files],
    }, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"pptx": str(pptx), "pdf": str(pptx.with_suffix('.pdf')),
                      "screenshot": str(screenshot), "slide_count": 10}, indent=2))


if __name__ == "__main__":
    main()
