"""Targeted GUI documentation update; preserve all other thesis evidence."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib,importlib.util,json
from docx import Document

B=Path(__file__).resolve().parent
ROOT=B.parents[1]
OLD=ROOT/'results/upright_v4_finalization_20260912_103145'

def sha(p):
    with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def main():
    state=json.loads((B/'initial_state.json').read_text());source=Path(state['backup'])
    assert sha(source)==state['input_word_sha256']
    doc=Document(source);changes=[];additions=[]
    spec=importlib.util.spec_from_file_location('preserved_styles',OLD/'build_thesis.py')
    h=importlib.util.module_from_spec(spec);spec.loader.exec_module(h)
    def match(prefix):
        ps=[p for p in doc.paragraphs if p.text.startswith(prefix)]
        assert len(ps)==1,(prefix,len(ps));return ps[0]
    def replace(prefix,new,reason):
        p=match(prefix);before=p.text;p.text=new
        changes.append({'kind':'paragraph','before':before,'after':new,'reason':reason})
    def insert(anchor,text,style=None):
        p=anchor.insert_paragraph_before(text,style=style)
        additions.append({'kind':'paragraph','text':text,'anchor':anchor.text});return p

    p=match('Der Ressourcenvergleich verwendet je Methode einen eigenen Prozess.')
    replace(p.text,p.text.replace('Eine grafische Oberfläche ist im nachgewiesenen Kern nicht enthalten. Der vorhandene frühere GUI-Prototyp ist damit keine validierte Bedienoberfläche des neuen Pakets.',
        'Die Messläufe enthalten keine GUI. Implementierung und Einbindung der vorhandenen Oberfläche werden in Abschnitt 6.8 getrennt beschrieben.'),
        'Separate measured headless runtime path from actually implemented GUI; explain integration in the dedicated subsection.')

    anchor=match('Evaluation und Ergebnisse')
    heading=insert(anchor,'Grafische Benutzeroberfläche und Integrationsstand','Heading 2')
    number=match('Festgelegte Versuche und Nachweisgrenzen')._p.pPr.numPr.numId.val
    h.set_num(heading,level=1,num=number)
    insert(anchor,
        'Die Oberfläche in src/live_tflite_gui.py ist als profilgebundener Entwicklungsstand umgesetzt (Belege: B11). Tkinter und Matplotlib zeigen Profilkennung, Rekonstruktionsfehler, Schwelle, Entscheidung, Zeitstempel sowie Angaben zu Erfassung, Inferenz und Pufferung. Der MSE-Verlauf enthält höchstens 100 Fenster; ungültige und veraltete Entscheidungen erhalten die Zustände ERROR beziehungsweise STALE. Start, Stop und Exit steuern die Messung, nicht den Lüfter. Ein Hintergrundthread führt die Auswertung aus; der Hauptthread aktualisiert die Anzeige. Eine begrenzte Nachrichtenqueue verhindert, dass eine verzögerte Anzeige den Auswerter beim Übergeben von Meldungen blockiert. Verworfene Anzeigemeldungen werden getrennt gezählt.')
    insert(anchor,
        'Die GUI verwendet live_tflite_monitor.py und live_pipeline.BufferedAcquisition mit der bereits korrigierten FIFO-Erfassung frischer Sensorwerte. Sie verarbeitet vollständige 128er-XYZ-Fenster fortlaufend, ohne die an einen Stellbefehl gebundene Abschnittsauswahl [180,300) s des finalen Tests. Ausgewertet wird der TFLite-Autoencoder eines per Kommandozeile gewählten Profils; RMS und Isolation Forest sind nicht eingebunden. scale_window wandelt Rohachsen zuerst in Float32 um und wendet den Profil-Scaler an. Die bei v4 vorgeschriebene achsenweise Mittelwertentfernung je Fenster in Float64 fehlt. Gleiche Tensorformen von 1 × 128 × 3 belegen daher keine gleiche Modelleingabe.')
    insert(anchor,
        'Die Profilauflösung verlangt profile.json mit Status ready sowie dazu passende Modell-, Scaler- und Schwellenartefakte. Das finale pilot_bundle.json besitzt ein anderes Schema und wird von der GUI nicht geladen. Alle vier vorhandenen Altprofile nennen 500 Hz; die voreingestellte Sensorerfassung liegt bei 200 Hz. Ohne Entwicklungsübersteuerung weist die Konsistenzprüfung diese Kombination zurück. Eine solche Übersteuerung oder ein bloßer Dateiaustausch stellt die v4-Vorverarbeitung nicht her. Die optionale Kommandozeilenüberschreibung der GUI-Schwelle ist ebenfalls keine zulässige Einstellung des eingefrorenen Methodenvergleichs.')
    cap=insert(anchor,'Tabelle 6-7: Implementierung und Nachweisgrenzen der vorhandenen GUI','Tabellenbeschriftung')
    rows=[
        ['Bereich','Umgesetzter Stand','Grenze für das finale Paket'],
        ['Anzeige und Bedienung','Status, MSE-Verlauf, Schwelle, Zeit-/Pufferangaben; Start, Stop, Exit.','Keine Lüftersteuerung; kein Drei-Methoden-Vergleich in der Oberfläche.'],
        ['Daten- und Modellweg','Profil-Autoencoder über live_tflite_monitor; Rohachsen → Float32 → Profil-Scaler.','Keine Anbindung an den finalen FrozenEngine-/WindowAssembler-Weg; Fensterzentrierung fehlt.'],
        ['Protokollierung','Monitor: Rohdaten, Entscheidungen und Ressourcen; GUI: Matplotlib-Zeichenabschluss und Meldungsverluste.','Kein Nachweis des physikalischen Anzeigezeitpunkts oder einer vergleichbaren Messreihe.'],
        ['Softwareprüfung','Simulationsmodus; Tests für Queue, Kontrollmeldungen, Journal und Veralterung.','Keine Aussage über tatsächliche Bildschirmanzeige, Sensor-Livebetrieb oder GUI-Zusatzlast mit v4.']
    ]
    t=doc.add_table(rows=1,cols=3)
    for c,v in zip(t.rows[0].cells,rows[0]):c.text=v
    for row in rows[1:]:
        for c,v in zip(t.add_row().cells,row):c.text=v
    h.style_table(t);anchor._p.addprevious(t._tbl)
    additions.append({'kind':'table','caption':cap.text,'rows':rows,'anchor':anchor.text})
    anchor=match('Erfüllungsgrad der Anforderungen')
    anchor.paragraph_format.page_break_before=True
    insert(anchor,
        'Für die GUI liegen 16 historische CSV-Dateien vom 23. bis 26. August 2026 mit insgesamt 30.075 Entscheidungszeilen vor. Sie sind den damaligen Profilen home_v001, fan_50 und fan_25 zugeordnet und enthalten sieben Felder einschließlich Score, Schwelle, Entscheidung, Zeitangabe und beobachteter Rate. Diese Dateien dokumentieren Ausgaben des damaligen GUI-Wegs. Sie enthalten jedoch keine ausreichende Zuordnung zum aktuellen Quellstand oder zum v4-Paket und keinen vollständigen CPU-, RAM- oder Anzeigezeitvergleich. Ein Abgleich tatsächlich dargestellter Bildschirmwerte mit den CSV-Werten ist dadurch nicht belegt.')
    insert(anchor,
        'Die ergänzende Prüfung vom 13. September 2026 ist eine Softwareprüfung ohne Hardware: Drei vorhandene GUI-Regressionen prüfen begrenzte Queues, erhaltene Kontrollmeldungen und Journalanlage sowie STALE bei alten Fenstern trotz frischer Entscheidung. Ein separater synthetischer Funktionstest bestätigt die unterschiedliche Mittelwertbehandlung des GUI- und v4-Wegs; er verwendet keine gemessenen Beschleunigungen und führt keine Modellinferenz aus. Diese Prüfungen ersetzen weder einen sichtbaren GUI-Funktionstest noch einen Sensorversuch. Die 18 Laufzeitprozesse aus Abschnitt 7.7 liefen ohne Oberfläche. Der in Abschnitt 4.3 geforderte Vergleich mit und ohne GUI bei ansonsten gleichen Bedingungen ist nicht erbracht; die vorhandenen Laufzeitwerte dürfen nicht als GUI-Messwerte verwendet werden.')

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells)==3 and row.cells[0].text=='A12':
                before=[c.text for c in row.cells]
                row.cells[1].text='GUI umgesetzt; Nachweis unvollständig'
                row.cells[2].text='Profilbasierte Oberfläche und begrenzte Softwaretests vorhanden. Finale v4-Anbindung, Bildschirm-/Ergebnisabgleich und Vergleich mit/ohne GUI nicht nachgewiesen.'
                changes.append({'kind':'requirement_row','before':before,'after':[c.text for c in row.cells],'reason':'Implemented optional feature still lacks the original conditional validation requirements; do not classify as absent or exempt.'})
                h.style_table(table)

    p=match('Die Latenzmessung beginnt beim hostseitig vollständigen Fenster;')
    replace(p.text,p.text.replace('Die optionale GUI ist im abschließenden Messweg nicht enthalten; ihre Zusatzlast mit dem neuen Paket wurde nicht geprüft.',
        'Die GUI ist umgesetzt, verwendet jedoch einen abweichenden Profilweg. Ihre Existenz und die hardwarefreien Teilprüfungen erfüllen A12 nicht vollständig; insbesondere bleibt der geforderte Vergleich der Zusatzlast offen.'),
        'Discuss implementation and incomplete A12 evidence without changing latency endpoints or other limitations.')
    replace('Die GUI bleibt eine optionale Erweiterung.',
        'Für eine spätere v4-Einbindung müsste die bestehende Anzeige ihre Entscheidungen aus dem finalen Datenweg erhalten, einschließlich eingefrorener Vorverarbeitung, Artefaktprüfung, Schwellen und Gültigkeitskennzeichnung. Der vorhandene Profil-Auswerter darf dabei nicht stillschweigend weiterverwendet werden. Anschließend wären angezeigte Werte gegen gespeicherte Fensterentscheidungen sowie Fehler-, Veraltungs- und Beendigungsfälle zu prüfen. Der ursprüngliche Nachweis verlangt außerdem einen Vergleich mit und ohne GUI bei ansonsten gleichen Eingaben, Methoden und Bedingungen: Verarbeitung bis zur Entscheidung und bis zum protokollierten Zeichenabschluss, CPU, RAM, Pufferverluste und ungültige Ausgaben. Offline-Replay und Sensor-Livebetrieb müssten getrennt ausgewiesen werden. Diese Integrations- und Messschritte sind nicht durchgeführt; der geprüfte Betrieb ohne Oberfläche bleibt davon getrennt.',
        'Name missing integration and original GUI comparison explicitly, without implementing it or waiving requirements.')

    for table in doc.tables:
        if table.rows[0].cells[0].text=='ID' and table.rows[0].cells[-1].text=='Projektpfad':
            before=[[c.text for c in row.cells] for row in table.rows]
            values=['B11','Gezielte GUI-Dokumentation: Quell-/Profilprüfung, historische Nachweise, hardwarefreie Tests und Word-Änderungsnachweis',str(B.relative_to(ROOT))+'/']
            for c,v in zip(table.add_row().cells,values):c.text=v
            h.style_table(table,is_appendix=True)
            changes.append({'kind':'whole_table','before':before,'after':[[c.text for c in row.cells] for row in table.rows],'reason':'Add separate GUI audit evidence without editing historical artifacts.'})
            break
    else:raise AssertionError('Digital evidence table missing')

    doc.core_properties.modified=datetime.now(timezone.utc)
    doc.core_properties.last_modified_by='Codex – gezielte GUI-Dokumentation'
    doc.core_properties.revision=(doc.core_properties.revision or 0)+1
    target=B/'thesis_gui_unformatted.docx';doc.save(target)
    manifest={'utc':datetime.now(timezone.utc).isoformat(),'input_sha256':sha(source),'draft_sha256':sha(target),
        'changes':changes,'additions':additions,'layout_changes':[{'anchor':'Erfüllungsgrad der Anforderungen','page_break_before':True,'reason':'Keep requirements table and following FF/H table legible without splitting the compact FF/H table.'}],'table_count':len(doc.tables),'scope':'GUI documentation only',
        'new_hardware_measurements':False,'model_or_threshold_changes':False,'new_gui_integration':False,
        'original_requirements_research_questions_conclusions_preserved':True}
    (B/'word_changes.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'draft':str(target),'changes':len(changes),'additions':len(additions),'tables':len(doc.tables)}))

if __name__=='__main__':main()
