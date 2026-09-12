#!/usr/bin/env python3
"""Build the final manuscript from the preserved source and explicit evidence.
Run with isolated python-docx environment; never modifies source data/models.
"""
from pathlib import Path
from copy import deepcopy
import argparse
import hashlib
import json
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

ROOT=Path('/home/malik/masterarbeit-edge-ai')
BASE=Path(__file__).resolve().parent
BACKUP=ROOT/'docs/backups/Akz_Masterarbeit_Bericht(3)_vor_abschluss_20260912_103145.docx'
SOURCE_SHA='a9a5d7d17b28098f554240a3e1650881e252db679f33319a1d11164b3246cbb8'
W= '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def digest(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def el(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e


def set_num(p,level=0,num=0):
 pp=p._p.get_or_add_pPr()
 for x in list(pp):
  if x.tag in (qn('w:numPr'),qn('w:ind')):pp.remove(x)
 n=el('w:numPr');n.append(el('w:ilvl',val=level));n.append(el('w:numId',val=num));pp.append(n)


def bookmark(p,name,number):
 p._p.insert(0,el('w:bookmarkStart',id=number,name=name));p._p.append(el('w:bookmarkEnd',id=number))


def field(p,instruction,cached='1'):
 r=el('w:r');r.append(el('w:fldChar',fldCharType='begin'));p._p.append(r)
 r=el('w:r');t=el('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';r.append(t);p._p.append(r)
 r=el('w:r');r.append(el('w:fldChar',fldCharType='separate'));p._p.append(r)
 p.add_run(cached)
 r=el('w:r');r.append(el('w:fldChar',fldCharType='end'));p._p.append(r)


def hyperlink(p,text,url):
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 r_id=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),r_id)
 r=el('w:r');pr=el('w:rPr');pr.append(el('w:color',val='1F4E79'));r.append(pr)
 t=el('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)


def text_runs(p,text):
 # Literal math/underscores are retained; only explicit bold/backtick formatting.
 for i,t in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)):
  if t.startswith('**') and t.endswith('**'):p.add_run(t[2:-2]).bold=True
  elif t.startswith('`') and t.endswith('`'):
   r=p.add_run(t[1:-1]);r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
  else:p.add_run(t)


def new_style(doc,name,base='Normal',size=12,bold=False):
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.base_style=doc.styles[base];s.font.name='Calibri';s.font.size=Pt(size);s.font.bold=bold
 s.font.color.rgb=RGBColor.from_string('000000')
 return s


def configure(doc):
 normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(12)
 normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
 normal.paragraph_format.widow_control=True
 normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 normal.element.get_or_add_rPr().append(el('w:lang',val='de-DE'))
 for level,size in ((1,16),(2,14),(3,12)):
  s=doc.styles[f'Heading {level}'];s.font.name='Calibri';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
  s.paragraph_format.space_before=Pt(15 if level>1 else 0);s.paragraph_format.space_after=Pt(8)
  s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
  s.paragraph_format.left_indent=Cm(0);s.paragraph_format.first_line_indent=Cm(0)
  s.paragraph_format.line_spacing=1.1
  s.paragraph_format.page_break_before=(level==1)
 for name in ['Tabellenbeschriftung','Abbildungsbeschriftung']:
  s=new_style(doc,name,'Caption',9.5);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(4)
  s.paragraph_format.keep_with_next=(name=='Tabellenbeschriftung');s.paragraph_format.keep_together=True
  s.paragraph_format.line_spacing=1.0
 s=new_style(doc,'Verzeichnisüberschrift','Heading 1',16,True)
 s.element.get_or_add_pPr().append(el('w:outlineLvl',val=9))
 s=new_style(doc,'FrontMatterHeading','Heading 1',16,True)
 s.element.get_or_add_pPr().append(el('w:outlineLvl',val=0))
 s=new_style(doc,'LiteratureEntry','Normal',10.5)
 s.paragraph_format.left_indent=Cm(.5);s.paragraph_format.first_line_indent=Cm(-.5);s.paragraph_format.line_spacing=1.05
 s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_together=True
 s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 for name in ('TOC 1','TOC 2','TOC 3','Table of Figures'):
  s=new_style(doc,name,'Normal',10.5)
  s.paragraph_format.line_spacing=1.0;s.paragraph_format.space_after=Pt(3)
  s.paragraph_format.keep_with_next=False;s.paragraph_format.keep_together=True


def heading_numbering(doc):
 # A single new multilevel list for ALL main headings fixes old mixed numIds.
 root=doc.part.numbering_part.element
 aid=max(int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum')))+1
 nid=max(int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num')))+1
 a=el('w:abstractNum',abstractNumId=aid);a.append(el('w:multiLevelType',val='multilevel'))
 for level in range(3):
  l=el('w:lvl',ilvl=level);l.append(el('w:start',val=1));l.append(el('w:numFmt',val='decimal'))
  l.append(el('w:lvlText',val='.'.join('%'+str(i+1) for i in range(level+1))))
  l.append(el('w:lvlJc',val='left'));l.append(el('w:suff',val='space'))
  p=el('w:pPr');p.append(el('w:ind',left=0,hanging=0));l.append(p);a.append(l)
 root.append(a);n=el('w:num',numId=nid);n.append(el('w:abstractNumId',val=aid));root.append(n)
 return nid


def style_table(t,is_appendix=False):
 t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 n=len(t.columns);width=16
 if is_appendix:widths=[1.4,5.0,9.6]
 elif n==2:widths=[4.0,12.0]
 elif n==3:widths=[3.5,4.7,7.8]
 elif n==4:widths=[3.0,3.7,3.7,5.6]
 elif n==5:widths=[4.0,2.5,3.0,3.2,3.3]
 elif n==6:widths=[4.0,2.4,2.4,2.4,2.4,2.4]
 else:widths=[width/n]*n
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 pr=t._tbl.tblPr
 for old in list(pr):
  if old.tag==qn('w:tblW'):pr.remove(old)
 pr.append(el('w:tblW',w=round(width/2.54*1440),type='dxa'))
 for j,row in enumerate(t.rows):
  trp=row._tr.get_or_add_trPr()
  if trp.find(qn('w:cantSplit')) is None:trp.append(el('w:cantSplit'))
  if j==0 and trp.find(qn('w:tblHeader')) is None:trp.append(el('w:tblHeader'))
  for k,c in enumerate(row.cells):
   c.width=Cm(widths[min(k,n-1)]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if j==0:
    cp=c._tc.get_or_add_tcPr();cp.append(el('w:shd',fill='E7EDF2',val='clear'))
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=(j==0);p.paragraph_format.keep_together=True
    p.alignment=(WD_ALIGN_PARAGRAPH.RIGHT if j>0 and k>0 and re.fullmatch(r'[−+\d][\d.,/% ×–−+()\[\] ]*',p.text) else WD_ALIGN_PARAGRAPH.LEFT)
    for r in p.runs:r.font.name='Calibri';r.font.size=Pt(9.5 if n<=5 else 9);r.bold=(j==0)


def runtime_text():
 a=json.loads((BASE/'runtime_analysis/analysis.json').read_text());rs=a['trials']
 def count(v):return format(v,',').replace(',','.')
 def fmt(v,n=3):return f'{v:.{n}f}'.replace('.',',')
 def span(values,n=3):return fmt(min(values),n)+'–'+fmt(max(values),n)
 names={'rms':'RMS','isolation_forest':'Isolation Forest','tflite_autoencoder':'TFLite-Autoencoder'}
 live=[r for r in rs if r['mode']=='sensor_live'];replay=[r for r in rs if r['mode']=='offline_replay']
 lines=['Für jede Methode wurden drei neue Prozesse im Sensor-Livebetrieb und drei getrennte Prozesse mit zeitgetreuem Offline-Replay ausgeführt. Die Reihenfolge der Methoden wurde zwischen den Wiederholungen rotiert. Jeder Prozess erfasst beziehungsweise spielt 300 s ab; die Modellverarbeitung bewertet ausschließlich den festgelegten Abschnitt [180,300) s. Daraus folgen je Prozess 120 s aktive Bewertungszeit. Zwanzig synthetische Aufwärmaufrufe erfolgten vor dem Messbeginn. Sie verändern keine Parameter.',
 'Jeder Sensorstart verwendet 75 % PWM und 25 kHz nach mindestens 60 s zusätzlicher softwareseitiger Auszeit. Ein Prozess enthält jeweils nur eine der drei Methoden. Alle neun Replay-Prozesse verwenden dieselbe vorhandene normale Rohaufnahme mit ihren ursprünglichen Host-Leseabständen. Diese Wiederholungen messen Softwareverhalten; sie sind keine zusätzlichen unabhängigen physikalischen Versuche.',
 'Tabelle 7-7: Latenz je Methode und Betriebsart; Bereich der drei getrennten Prozesse',
 '| Betriebsart / Methode | P99 gesamt [ms] | P99 Rechenkern [ms] | größtes Maximum gesamt [ms] | Überschreitungen |',
 '|---|---:|---:|---:|---:|']
 for mode,title in [('sensor_live','Sensor'),('offline_replay','Replay')]:
  for method,name in names.items():
   group=[r for r in rs if r['mode']==mode and r['method']==method]
   lines.append('| '+title+' / '+name+' | '+span([r['latency_ms']['complete_to_decision']['p99'] for r in group])+' | '+span([r['latency_ms']['model_core']['p99'] for r in group])+' | '+fmt(max(r['latency_ms']['complete_to_decision']['maximum'] for r in group))+' | '+str(sum(r['deadline_exceedances'] for r in group))+' |')
 lines+=['Die gesamte Verarbeitungslatenz beginnt beim Host-Leseabschluss des letzten XYZ-Punkts eines Fensters und endet bei der verfügbaren Entscheidung vor dem Schreiben ins Journal. Sie enthält Rohdatenübergabe, Fensteraufbau, Qualitätsprüfung, Mittelwertentfernung, Queue, Standardisierung und Auswertung. Die Sensorwandlungszeit vor der Host-Lesung sowie die Fensterfüllzeit sind nicht enthalten. Der Rechenkern bezeichnet die RMS-Berechnung, IsolationForest.score_samples beziehungsweise nur Interpreter.invoke; Tensorübertragung und Autoencoder-Fehlerberechnung sind gesonderte Anteile.',
 f"Das beobachtete Fensterintervall beträgt im Sensorbetrieb {span([r['window_deadline_ms'] for r in live])} ms. Das empirische P99 liegt {'in allen neun Prozessen darunter' if all(r['latency_ms']['complete_to_decision']['p99']<r['window_deadline_ms'] for r in live) else 'nicht in allen Prozessen darunter'}. Unter {count(sum(r['valid_decisions'] for r in live))} gültigen Sensorentscheidungen wurden {sum(r['deadline_exceedances'] for r in live)} Fristüberschreitungen gezählt. H2 ist damit {'im festgelegten Messumfang bestätigt' if all(r['latency_ms']['complete_to_decision']['p99']<r['window_deadline_ms'] for r in live) else 'im Messumfang nicht durchgängig bestätigt'}. Ein P99 ist keine Zusage für jede zukünftige Entscheidung und kein Beleg harter Echtzeit.",
 f'![Abbildung 7-5: Verarbeitungslatenz aus drei eigenen Prozessen je Methode und Betriebsart. Punkte zeigen P99, Linien den Bereich vom Median zum Maximum; die gestrichelte Linie kennzeichnet das kleinste beobachtete Fensterintervall. Quelle: eigene Laufzeitmessungen.]({BASE}/runtime_analysis/latency_comparison.png)',
 'Tabelle 7-8: Prozessressourcen im aktiven Abschnitt [180,300) s; Bereich der drei Prozesse',
 '| Betriebsart / Methode | mittlere CPU [% eines Kerns] | maximales RSS [MiB] | gültig / ungültig gesamt | verworfene Fenster |',
 '|---|---:|---:|---:|---:|']
 for mode,title in [('sensor_live','Sensor'),('offline_replay','Replay')]:
  for method,name in names.items():
   group=[r for r in rs if r['mode']==mode and r['method']==method]
   lines.append('| '+title+' / '+name+' | '+span([r['resource_primary_180_300']['process_cpu_percent_one_core_time_weighted'] for r in group],2)+' | '+span([r['resource_primary_180_300']['process_rss_mib']['maximum'] for r in group],2)+' | '+str(sum(r['valid_decisions'] for r in group))+' / '+str(sum(r['invalid_decisions'] for r in group))+' | '+str(sum(r['pipeline']['windows_dropped'] for r in group))+' |')
 lines+=['Die CPU-Auslastung ist aus den ungefähr 100-ms-Prozessmessungen zeitgewichtet; 100 % entsprechen einem logischen Kern. RSS enthält den gesamten jeweiligen Python-Prozess einschließlich Bibliotheken, Vorverarbeitung, Erfassung und Protokollierung. Beim Replay liegt die komplette Quelldatei bereits vor Beginn in einem DataFrame und einer Liste von Datensätzen im Speicher. Dieser zusätzliche Speicher gehört zum Replay-Verfahren und darf nicht als reiner Modellbedarf oder als identisch zum Sensorbetrieb interpretiert werden. Der kontinuierliche Monitor erfasst die Aufnahme beziehungsweise Wiedergabe; Modellstart und abschließendes erneutes Laden der Sensor-CSV werden außerhalb dieses Ressourcenverlaufs durch Ladezeiten und Umgebungssnapshots dokumentiert. Für CPU und RAM ist der Prozess die Vergleichseinheit. Die erste CPU-Probe besitzt keinen vorherigen Monitorzeitstempel und ist deshalb aus den zeitgewichteten Kennwerten und der CPU-Grafik ausgeschlossen; der ursprüngliche Wert bleibt in der Rohdatei erhalten. Ein sehr kurzer erster Messabstand kann einen überhöhten Quotienten aus CPU-Zeit und Wandzeit erzeugen.',
 f'![Abbildung 7-6: CPU- und RSS-Verläufe im Sensor-Livebetrieb. Zur Darstellung wurden die CPU-Werte in 1-s-Abschnitten zeitgewichtet und die RSS-Werte gemittelt; die Kennwerte der Tabelle verwenden die ursprünglichen zeitgewichteten Messintervalle. Quelle: eigene Laufzeitmessungen.]({BASE}/runtime_analysis/sensor_live_resources.png)',
 'Die Rohreihen, maximalen Host-Abstände, Speichertrends, Modellladezeiten, Aufwärmzeiten, Übertragungsanteile und Latenzperzentile je vollständigem Lauf sind im digitalen Laufzeitbericht erhalten. Die Host-Abstände an den Grenzen aufeinanderfolgender vollständiger Fenster sowie die Intervalle ihrer Ankunft werden ebenfalls separat ausgewiesen; sie sind keine direkte Messung von Sensorwandlungspausen. Die neue Instrumentierung schreibt Entscheidungen und Ressourcen fortlaufend und begrenzt die Queue auf vier Fenster. Rohdaten werden bei einem verworfenen Entscheidungsfenster weiterhin gespeichert.',
 f"Die neun Sensoraufnahmen enthalten insgesamt {count(sum(r['quality']['xyz_points'] for r in live))} XYZ-Punkte bei {span([r['quality']['observed_xyz_per_second'] for r in live],6)} XYZ/s. Der erste Punkt liegt {span([r['quality']['first_xyz_after_command_s']*1000 for r in live])} ms nach dem Stellbefehl. Die maximale beobachtete Host-Lesepause über alle Sensorläufe beträgt {fmt(max(r['quality']['host_interval_max_ms'] for r in live))} ms. Es wurden {sum(r['quality']['gap_flagged_xyz'] for r in live)} Gap-, {sum(r['quality']['overrun_flagged_xyz'] for r in live)} Overrun- und {sum(r['quality']['saturated_flagged_xyz'] for r in live)} Sättigungsflags festgestellt. Diese Nullbefunde ersetzen keinen Nachweis unbekannter physischer Verluste.",
 f"Über beide Betriebsarten stimmen alle {count(sum(r['scores_compared'] for r in rs))} ausgegebenen gültigen Scores exakt mit der erneuten eingefrorenen Offline-Auswertung derselben Rohfenster überein. Die maximale absolute Abweichung beträgt {fmt(max(r['source_to_reference_score_max_abs_difference'] for r in rs),1)}. Auch die Entscheidungen stimmen überein. Die Rückstandsprüfung findet {sum(r['pipeline']['unprocessed_windows'] for r in rs)} vollständig gebildete, nach Abschluss unverarbeitete Fenster. Die numerische Prüfung beschreibt Konsistenz und ist kein zusätzlicher Erkennungsnachweis."]
 deltas=[r['resource_primary_180_300']['rss_first_last_change_mib'] for r in live]
 lines += [f"Zwischen erster und letzter Monitorprobe im aktiven Sensorabschnitt steigt beziehungsweise verändert sich der RSS um {span(deltas)} MiB. Dieser Verlauf wird nicht als konstanter Speicherbedarf dargestellt. Das Ausbleiben eines Speicherabbruchs und eines wachsenden Fenster-Rückstands in 300 s ist kein unbegrenzter Dauerbetriebsnachweis. Eine Ursache des Speicherverlaufs lässt sich aus RSS allein nicht bestimmen."]
 lines += ['Eine getrennte explorative Speicherprobe führte sechsmal dieselbe gespeicherte Aufnahme durch den Fensteraufbau und den RMS-Auswerter. Sie enthielt weder Sensorzugriff noch Queue oder Journale. Nach 1.164 Auswertungen waren rund 443 KiB aktuelle Python-Allokationen registriert; nach einem expliziten GC-Aufruf waren es rund 89 KiB. Der RSS sank dabei nicht entsprechend. Diagnose-Snapshots verursachen selbst Speicherbedarf. Die Probe weist damit keine unbegrenzte Speicherung sämtlicher Fenster nach, erklärt aber auch nicht den gesamten Speicherverlauf des Sensorbetriebs. Sie ist kein weiterer Laufzeitbenchmark. Einzelwerte und Einschränkungen stehen im digitalen Speicherbericht unter B7.']
 sizes=a['artifact_sizes_bytes']
 lines+= [f"Das TFLite-Modell belegt {count(sizes['autoencoder_float32.tflite'])} Byte, der gespeicherte Isolation Forest {count(sizes['isolation_forest.joblib'])} Byte und der gemeinsame Scaler {count(sizes['scaler.joblib'])} Byte. Die Dateigrößen sind von den gemessenen Prozessspeichern zu unterscheiden.",
 'Softwareumgebung: Raspberry Pi 5 mit vier logischen Kernen, 64-Bit-Linux, Python 3.13.5, NumPy 2.5.1, pandas 3.0.3, scikit-learn 1.9.0, ai-edge-litert 2.1.6 und psutil 7.2.2. OMP_NUM_THREADS, OPENBLAS_NUM_THREADS und MKL_NUM_THREADS waren auf 1 gesetzt. Die Umgebung und Temperatur-/Drosselungsrücklesungen wurden vor und nach jedem Prozess protokolliert. Es lief keine GUI. Kurze Statusabfragen und leichte Dokumentvorbereitung erfolgten im Hintergrund; absichtliche Trainings-, Test- und Renderinglast war während der Messmatrix ausgesetzt.',
 'Nach jedem Sensorlauf und abschließend wurde 0 % PWM zurückgelesen. Dies dokumentiert den Steuerungszustand, nicht einen neuen mechanischen Sichtnachweis. Die längeren gesamten Auszeiten ergeben sich aus vorherigem Stopp, Abschluss und Vorbereitung des nächsten Prozesses; identisch sind nur die vorgesehenen zusätzlichen 60 s. Für einen unbegrenzten Dauerbetrieb, GUI-Zusatzlast oder einen Ausfall des Rechners wird kein Nachweis behauptet.']
 return '\n\n'.join(lines)+'\n'


def bibliography_text():
 b=json.loads((BASE/'bibliography.json').read_text())
 return '\n\n'.join(x['text']+' '+x['url']+(' '+x['additional_url'] if x.get('additional_url') else '')+' (Abruf: 12.09.2026.)' for x in b['entries'])


def build(output,draft=False):
 assert digest(BACKUP)==SOURCE_SHA
 doc=Document(BACKUP);pars=list(doc.paragraphs);body=doc._element.body
 # Preserve actual main text and tables of chapters 1–5, with original bookmarks.
 first_main=pars[60]._p;first_replace=pars[231]._p
 main_source=[]
 for e in list(body)[list(body).index(first_main):list(body).index(first_replace)]:
  main_source.extend(t.text or '' for t in e.iter(qn('w:t')))
 main_section=deepcopy(pars[295]._p.pPr.sectPr)
 front_section=deepcopy(pars[59]._p.pPr.sectPr)
 abbreviations=deepcopy(doc.tables[0]._tbl)
 # Replace previous contents/caption lists with updatable fields.
 i=list(body).index(pars[32]._p);j=list(body).index(first_main)
 for x in list(body)[i:j]:body.remove(x)
 # Remove a redundant empty continuous section after the cover.
 body.remove(pars[19]._p)
 # Replace chapter 6 through bibliography/appendix; retain chapters 1–5.
 for x in list(body)[list(body).index(first_replace):]:body.remove(x)
 body.append(main_section)
 configure(doc);num_id=heading_numbering(doc)
 # Direct formatting overrides old heading list ids and inconsistent type sizes.
 main_active=False
 for p in doc.paragraphs:
  if p._p is first_main:main_active=True
  if p.style.name.startswith('Heading'):
   if main_active:
    level=int(p.style.name.rsplit(' ',1)[1])-1;set_num(p,level,num_id)
   else:set_num(p,0,0);p.paragraph_format.page_break_before=False
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT
   for r in p.runs:r.font.size=None;r.font.color.rgb=RGBColor(0,0,0)
  if main_active and re.match(r'^Tabelle \d+-\d+:',p.text):p.style='Tabellenbeschriftung'
 # Insert front matter before chapter 1, using a small temporary document section.
 def insert_p(text='',style=None):
  p=doc.add_paragraph(text,style);body.remove(p._p);first_main.addprevious(p._p);return p
 front=[]
 for title,code in [('Inhaltsverzeichnis','TOC \\o "1-3" \\h \\z \\u'),
                    ('Abbildungsverzeichnis','TOC \\t "Abbildungsbeschriftung,1" \\h \\z'),
                    ('Tabellenverzeichnis','TOC \\t "Tabellenbeschriftung,1" \\h \\z')]:
  h=insert_p(title,'Verzeichnisüberschrift' if title=='Inhaltsverzeichnis' else 'FrontMatterHeading');set_num(h,0,0)
  h.paragraph_format.page_break_before=True
  p=insert_p();field(p,code,'Verzeichnis wird bei der abschließenden Formatprüfung aktualisiert.');front.append(p)
 h=insert_p('Abkürzungsverzeichnis','FrontMatterHeading');set_num(h,0,0);h.paragraph_format.page_break_before=True
 first_main.addprevious(abbreviations)
 p=insert_p();p._p.get_or_add_pPr().append(front_section)
 # Use decimal page numbers for the entire main text, appendix, and bibliography.
 pn=main_section.find(qn('w:pgNumType'))
 if pn is None:pn=el('w:pgNumType');main_section.append(pn)
 pn.set(qn('w:fmt'),'decimal');pn.set(qn('w:start'),'1')
 for section in doc.sections:
  section.left_margin=Cm(2.5);section.right_margin=Cm(2.5);section.top_margin=Cm(2.5);section.bottom_margin=Cm(2)
  section.header_distance=Cm(1);section.footer_distance=Cm(1)
 for i,section in enumerate(doc.sections):
  for kind in ('header','even_page_header','first_page_header','footer','even_page_footer','first_page_footer'):
   part=getattr(section,kind);part.is_linked_to_previous=False
   for e in list(part._element):part._element.remove(e)
   p=part.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   if 'footer' in kind and i>0:field(p,'PAGE','1')
 # Retain declaration wording, but do not carry an old signature image into
 # this substantially updated manuscript. The preserved source retains it.
 body.remove(pars[27]._p)
 pars[28].text='Unterschrift: __________________________'
 # Keep title/declaration; leave a blank for the author's actual signing date.
 for p in doc.paragraphs:
  if p.text=='Datum: Frankfurt, den XXXXX':
   p.text='Datum: Frankfurt, den __________________'
 md=(BASE/'chapters_completion.md').read_text()
 md=md.replace('{{RUNTIME_RESULTS}}','Noch nicht freigegebener Entwurf: Laufzeitmatrix läuft.' if draft else runtime_text())
 md=re.sub(r'(\|)\n\n(?=\|)',r'\1\n',md)
 md=md.replace('{{BIBLIOGRAPHY}}',bibliography_text())
 assert '{{' not in md
 (BASE/('manuscript_draft.md' if draft else 'manuscript_completed.md')).write_text(md)
 lines=md.splitlines();i=0;appendix=False;literature=False;pending_caption=False
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('# '):
   title=line[2:];appendix=title=='Anhang';literature=title=='Literaturverzeichnis'
   p=doc.add_heading(title,1);set_num(p,0,0 if appendix or literature else num_id);continue
  if line.startswith('## '):
   p=doc.add_heading(line[3:],2);set_num(p,1,0 if appendix else num_id);continue
  if line.startswith('### '):
   p=doc.add_heading(line[4:],3);set_num(p,2,0 if appendix else num_id);continue
  m=re.match(r'^!\[(.*?)\]\((.*)\)$',line)
  if m:
   path=Path(m[2]);assert path.exists(),path
   # Dedicated publication-size rendering uses the same archived values.
   if not draft:
    mapping={'Abbildung 7-1':'plate_scores','Abbildung 7-2':'plate_rms','Abbildung 7-3':'axis_decomposition','Abbildung 7-4':'false_alarm_comparison','Abbildung 7-5':'latency_comparison','Abbildung 7-6':'sensor_live_resources'}
    figure_id=m[1].split(':',1)[0];assert figure_id in mapping,figure_id
    path=BASE/'word_figures'/(mapping[figure_id]+'.png');assert path.exists(),path
   import struct
   with path.open('rb') as im:
    header=im.read(24);assert header[:8]==b'\x89PNG\r\n\x1a\n'
    w,h=struct.unpack('>II',header[16:24])
   width=min(16,17.5*w/h)
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
   p.add_run().add_picture(str(path),width=Cm(width))
   cap=doc.add_paragraph(m[1],style='Abbildungsbeschriftung');cap.paragraph_format.keep_with_next=False
   continue
  if re.match(r'^Tabelle (\d+|A)-\d+:',line):
   doc.add_paragraph(line,'Tabellenbeschriftung');pending_caption=True;continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
   vals=[[x.strip() for x in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
   assert len(vals)>1 and len(set(map(len,vals)))==1,vals
   t=doc.add_table(rows=len(vals),cols=len(vals[0]))
   for row,vs in zip(t.rows,vals):
    for cell,value in zip(row.cells,vs):text_runs(cell.paragraphs[0],value)
   style_table(t,appendix and len(t.columns)==3);pending_caption=False
   doc.add_paragraph().paragraph_format.space_after=Pt(0)
   continue
  p=doc.add_paragraph(style='LiteratureEntry' if literature else None)
  if literature:
   for s in re.split(r'(https?://\S+)',line):
    if s.startswith('http'):hyperlink(p,s,s)
    else:p.add_run(s)
  else:text_runs(p,line)
 # Abbreviations are extended without changing the established terms.
 from docx.table import Table
 at=Table(abbreviations,doc._body)
 existing={r.cells[0].text for r in at.rows}
 for key,value in [('AC','Wechselanteil nach Entfernung des Gleichanteils'),('IF','Isolation Forest'),('LSB','Least Significant Bit'),('MiB','Mebibyte (2²⁰ Byte)'),('P99','Empirisches 99. Perzentil'),('PSD','Power Spectral Density (Leistungsdichtespektrum)')]:
  if key not in existing:
   cells=at.add_row().cells;cells[0].text=key;cells[1].text=value
 # Format all retained and new tables; retain their text verbatim.
 for t in doc.tables:
  style_table(t,len(t.columns)==3 and any('B1'==r.cells[0].text for r in t.rows))
 # Explicit bookmarks support index refresh and stable cross references.
 toc=[];tables=[];figures=[];counts=[0,0,0];bid=10000
 main=False
 for p in doc.paragraphs:
  if p._p is first_main:main=True
  if not main:continue
  if p.style.name.startswith('Heading'):
   level=int(p.style.name.rsplit(' ',1)[1])-1
   n=p._p.pPr.numPr.numId.val
   if n==num_id:
    counts[level]+=1
    for j in range(level+1,3):counts[j]=0
    prefix='.'.join(str(x) for x in counts[:level+1])+' '
   else:prefix=''
   name=f'FinalHeading_{bid}';bookmark(p,name,bid);toc.append(dict(level=level,title=prefix+p.text,bookmark=name));bid+=1
  elif p.style.name in ('Tabellenbeschriftung','Abbildungsbeschriftung'):
   name=f'FinalCaption_{bid}';bookmark(p,name,bid);bid+=1
   (tables if p.style.name=='Tabellenbeschriftung' else figures).append(dict(title=p.text,bookmark=name))
 # Keep this manifest for content and page-number verification after rendering.
 actual_main=[]
 for e in list(body)[list(body).index(first_main):]:
  if e.tag==qn('w:p') and ''.join(t.text or '' for t in e.iter(qn('w:t')))=='Implementierung':break
  actual_main.extend(t.text or '' for t in e.iter(qn('w:t')))
 assert actual_main==main_source,'Chapter 1–5 text changed'
 settings=doc.settings.element
 old=settings.find(qn('w:updateFields'))
 if old is not None:settings.remove(old)
 settings.append(el('w:updateFields',val='true'))
 from datetime import datetime,timezone
 doc.core_properties.modified=datetime.now(timezone.utc)
 doc.core_properties.last_modified_by='Codex – technische Dokumentbearbeitung'
 doc.core_properties.revision=(doc.core_properties.revision or 0)+1
 output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
 manifest={'backup_sha256':SOURCE_SHA,'output':str(output),'draft':draft,'main_chapters_1_5_text_preserved':True,
           'headings':toc,'tables':tables,'figures':figures,'table_count':len(doc.tables),'paragraph_count':len(doc.paragraphs),
           'text_word_count':len(' '.join(p.text for p in doc.paragraphs).split()),'output_sha256':digest(output)}
 (BASE/('word_draft_manifest.json' if draft else 'word_build_manifest.json')).write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
 print(json.dumps({k:v for k,v in manifest.items() if k not in ['headings','tables','figures']},ensure_ascii=False,indent=2))


if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--draft',action='store_true');a=p.parse_args();build(a.output,a.draft)
