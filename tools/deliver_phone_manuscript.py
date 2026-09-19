#!/usr/bin/env python3
"""Install the reviewed binary report only after source/preservation checks pass."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import shutil
from docx import Document
from docx.table import Table
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/phone_vibration_word_revision_20260919'
NAME='Akz_Masterarbeit_Bericht(3)'
TARGET=ROOT/'docs'/f'{NAME}.docx'
BACKUP=ROOT/'docs/backups'/f'{NAME}_vor_handyversuch_20260919.docx'


def sha(p):
    with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def cells(table):return [[c.text for c in row.cells] for row in table.rows]


def bibliography(doc):
    paragraphs=doc.paragraphs
    start=next(i for i,p in enumerate(paragraphs) if p.text=='Literaturverzeichnis')
    return [p.text for p in paragraphs[start:]]


def main():
    manifest=json.loads((OUT/'revision_manifest.json').read_text())
    evidence=json.loads((OUT/'evidence_audit.json').read_text())
    layout=json.loads((OUT/'document_checks.json').read_text())
    final=OUT/f'{NAME}.docx';pdf=OUT/f'{NAME}.pdf'
    assert evidence['status']=='passed' and evidence['total_original_decisions_reproduced']==1205
    assert layout['automated_status']=='passed' and layout['docx_sha256']==sha(final) and layout['pdf_sha256']==sha(pdf)
    assert sha(BACKUP)==manifest['source_sha256']
    assert sha(TARGET)==sha(BACKUP),'Target changed since backup; do not overwrite a newer user edit.'
    for name,digest in evidence['source_sha256'].items():assert sha(ROOT/name)==digest,name
    for name,digest in manifest['figures']['source_sha256'].items():assert sha(ROOT/name)==digest,name
    old=Document(BACKUP);new=Document(final);blocks=list(old._element.body)
    new_tables=[cells(t) for t in new.tables]
    prior_tables=[e for e in blocks[:230] if e.tag==qn('w:tbl')]
    for e in prior_tables:assert cells(Table(e,old._body)) in new_tables
    for i in [266,359,366,431,434,440]:assert cells(Table(blocks[i],old._body)) in new_tables,i
    assert bibliography(old)==bibliography(new),'Bibliographic content changed'
    # The entire unchanged cover/front declaration precedes the summary heading.
    def front(d):
        pp=d.paragraphs;stop=next(i for i,p in enumerate(pp) if p.text=='Zusammenfassung')
        return [p.text for p in pp[:stop]]
    assert front(old)==front(new)
    assert len(new.sections)==3
    result={'completed_utc':datetime.now(timezone.utc).isoformat(),'source_sha256':sha(BACKUP),'pages':layout['pages'],'new_figures':9,'tables':len(new.tables),'original_decisions_reproduced':1205,'historical_tables_retained_exactly':True,'pre_chapter6_tables_retained_exactly':len(prior_tables),'bibliography_text_retained_exactly':True,'cover_and_declaration_retained_exactly':True,'raw_sources_unchanged':True,'visual_review':'Alle PDF-Seiten in Kontaktbögen, neue Diagramme und zentrale Ergebnis-/Anhangseiten geprüft; keine abgeschnittenen Diagramme, keine verwaisten Kapitelübergänge.','backup':str(BACKUP.relative_to(ROOT)),'outputs':{}}
    for source in [final,pdf]:
        target=ROOT/'docs'/source.name
        if source.suffix=='.pdf' and target.exists():
            backup_pdf=BACKUP.with_suffix('.pdf')
            assert not backup_pdf.exists(),'Existing PDF backup requires explicit review before overwrite.'
            shutil.copy2(target,backup_pdf)
            result['previous_pdf_backup']=str(backup_pdf.relative_to(ROOT))
        temporary=target.with_name(target.name+'.phone_revision.tmp')
        assert not temporary.exists()
        shutil.copy2(source,temporary)
        if source.suffix=='.docx':assert sha(TARGET)==sha(BACKUP)
        temporary.replace(target)
        assert sha(target)==sha(source)
        result['outputs'][str(target.relative_to(ROOT))]=sha(target)
    manifest['target_replaced']=True
    manifest['delivered_docx_sha256']=sha(TARGET)
    (OUT/'revision_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    (OUT/'delivery_manifest.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps(result,indent=2,ensure_ascii=False))


if __name__=='__main__':main()
