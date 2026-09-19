#!/usr/bin/env python3
"""Add evidence-based training/calibration figures to the reviewed thesis.

No training, calibration, hardware access or changes to experimental data.
The original DOCX is backed up; only a staged document is produced here.
Run with the project's venv and python-docx available on PYTHONPATH.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, MaxNLocator
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/masterarbeit_ueberarbeitet_revision_20260919'
SOURCE = ROOT / 'docs/Masterarbeit_ueberarbeitet.docx'
BACKUP = ROOT / 'docs/backups/Masterarbeit_ueberarbeitet_vor_diagrammen_20260919.docx'
EXPECTED_SOURCE_HASH = 'f55197f99e681eca49649913f519406b592695d6c84d26f7c7974514628dc050'
PROFILE = ROOT / 'profiles/standlauf_pwm75_200hz'
BUNDLE = ROOT / 'results/phone_vibration_pilot_20260919/bundle'
SOURCES = {
    'history': PROFILE / 'results/training_history.csv',
    'keras_errors': PROFILE / 'results/validation_reconstruction_errors.csv',
    'comparison_scores': BUNDLE / 'validation_predictions.csv',
    'bundle': BUNDLE / 'bundle.json',
}
COLORS = ['#007F86', '#B46A08', '#7753A2']


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(path):
    with path.open(newline='') as stream:
        return list(csv.DictReader(stream))


def de(value, decimals=3):
    return f'{value:.{decimals}f}'.replace('.', ',')


def make_figures():
    figures = OUT / 'figures'
    figures.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        'font.family': 'DejaVu Sans', 'font.size': 9,
        'axes.labelsize': 9, 'axes.titlesize': 10,
        'xtick.labelsize': 8, 'ytick.labelsize': 8,
        'legend.fontsize': 8.5, 'axes.spines.top': False,
        'axes.spines.right': False, 'axes.edgecolor': '#65717A',
        'pdf.fonttype': 42, 'savefig.facecolor': 'white',
    })
    history = rows(SOURCES['history'])
    epochs = np.array([int(r['epoch']) for r in history])
    training = np.array([float(r['loss']) for r in history])
    validation = np.array([float(r['val_loss']) for r in history])
    assert np.array_equal(epochs, np.arange(1, 101))
    assert epochs[np.argmin(validation)] == 100
    fig, ax = plt.subplots(figsize=(6.25, 3.2))
    ax.plot(epochs, training, color=COLORS[0], lw=1.8,
            label='Training (291 Fenster)')
    ax.plot(epochs, validation, color=COLORS[1], lw=1.8, ls='--',
            label='Validierung (97 Fenster)')
    ax.scatter([100, 100], [training[-1], validation[-1]],
               color=COLORS[:2], s=22, zorder=5)
    ax.axvline(100, color='#89939C', lw=.8, ls=':')
    ax.annotate(f'Validierung: {de(validation[-1], 5)}',
                xy=(100, validation[-1]), xytext=(60, .58),
                arrowprops={'arrowstyle': '-', 'color': COLORS[1]}, color=COLORS[1])
    ax.annotate(f'Training: {de(training[-1], 5)}',
                xy=(100, training[-1]), xytext=(61, .08),
                arrowprops={'arrowstyle': '-', 'color': COLORS[0]}, color=COLORS[0])
    ax.set(xlabel='Trainingsepoche', ylabel='Mittlerer quadratischer Fehler (MSE)',
           xlim=(1, 103), ylim=(0, 1.34))
    ax.set_xticks([1, 20, 40, 60, 80, 100])
    ax.yaxis.set_major_formatter(FuncFormatter(lambda x, _: de(x, 1)))
    ax.grid(axis='y', color='#DFE4E8', lw=.6)
    ax.legend(loc='upper right', frameon=False)
    fig.tight_layout(pad=1.0)
    for extension in ('png', 'pdf'):
        fig.savefig(figures / f'training_history.{extension}', dpi=300)
    plt.close(fig)

    comparison = rows(SOURCES['comparison_scores'])
    keras = rows(SOURCES['keras_errors'])
    ae_threshold = float(keras[0]['threshold'])
    ae = np.array([float(r['reconstruction_error_mse']) for r in keras])
    series = [('Autoencoder (Keras)', 'MSE', ae, ae_threshold)]
    for key, title, axis in [('isolation_forest', 'Isolation Forest', '−score_samples'),
                             ('rms', 'RMS', 'RMS standardisierter Werte')]:
        selected = [r for r in comparison if r['method'] == key]
        assert len({r['threshold'] for r in selected}) == 1
        values = np.array([float(r['score']) for r in selected])
        series.append((title, axis, values, float(selected[0]['threshold'])))
    metrics = {}
    fig, axes = plt.subplots(1, 3, figsize=(6.25, 3.25), sharey=True)
    for i, (ax, (title, label, scores, threshold)) in enumerate(zip(axes, series)):
        assert len(scores) == 97 and np.all(np.isfinite(scores))
        ordered = np.sort(scores)
        expected = np.quantile(scores, .99, method='linear')
        assert np.isclose(expected, threshold, rtol=0, atol=1e-12), (title, expected, threshold)
        above = int(np.sum(scores > threshold))
        assert above == 1
        span = ordered[-1] - ordered[0]
        left, right = ordered[0] - .08*span, ordered[-1] + .08*span
        xs = np.r_[left, ordered, right]
        ys = np.r_[0, np.arange(1, 98) / 97 * 100, 100]
        ax.step(xs, ys, where='post', color=COLORS[i], lw=1.5)
        ax.axvline(threshold, color='#AE3437', ls='--', lw=1.15)
        ax.scatter(ordered[-2:], np.array([96, 97])/97*100,
                   s=16, color=COLORS[i], zorder=5)
        ax.set_title(f'{title}\nSchwelle: {de(threshold, 6)}', fontsize=9)
        ax.set(xlabel=label, xlim=(left, right), ylim=(0, 106))
        ax.set_yticks([0, 25, 50, 75, 100])
        ax.xaxis.set_major_locator(MaxNLocator(nbins=3))
        ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: de(x, 2)))
        ax.grid(axis='y', color='#DFE4E8', lw=.6)
        metrics[title] = {
            'n': 97, 'threshold': threshold, 'linear_p99_recomputed': float(expected),
            'minimum': float(ordered[0]), 'maximum': float(ordered[-1]),
            'ordered_value_96': float(ordered[-2]),
            'above_threshold': above,
        }
    axes[0].set_ylabel('Anteil der Kalibrierungsfenster [%]')
    fig.text(.5, .025, 'Gestrichelt: eingefrorene P99-Schwelle · Punkte: die zwei größten Scores',
             ha='center', fontsize=8)
    fig.tight_layout(rect=(0, .08, 1, 1), pad=.9, w_pad=1.2)
    for extension in ('png', 'pdf'):
        fig.savefig(figures / f'calibration_scores.{extension}', dpi=300)
    plt.close(fig)
    return {
        'training': {'epochs': 100, 'best_validation_epoch': 100,
                     'final_loss': float(training[-1]), 'final_val_loss': float(validation[-1])},
        'calibration': metrics,
        'ae_calibration_source': 'Original Keras scores, not recomputed TFLite scores; frozen threshold unchanged.',
    }


def change_text(paragraph, text):
    """Preserve paragraph properties/bookmarks and the first run formatting."""
    nodes = paragraph._p.xpath('.//w:t')
    assert nodes, paragraph.text
    nodes[0].text = text
    for node in nodes[1:]:
        node.text = ''


def find_paragraph(document, prefix):
    found = [p for p in document.paragraphs if p.text.startswith(prefix)]
    assert len(found) == 1, (prefix, len(found))
    return found[0]


def add_figure_before(anchor, filename, caption, alt):
    picture = anchor.insert_paragraph_before()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.space_before = Pt(6)
    picture.paragraph_format.space_after = Pt(0)
    picture.paragraph_format.line_spacing = 1
    shape = picture.add_run().add_picture(str(OUT / 'figures' / filename), width=Cm(15.8))
    shape._inline.docPr.set('descr', alt)
    caption_paragraph = anchor.insert_paragraph_before(caption, style='Abbildungsbeschriftung')
    caption_paragraph.paragraph_format.keep_together = True
    caption_paragraph.paragraph_format.keep_with_next = False


def revise_document():
    doc = Document(BACKUP)
    original_tables = len(doc.tables)
    original_shapes = len(doc.inline_shapes)

    literature = find_paragraph(doc, 'Die hier herangezogene Preprintfassung von Garay')
    change_text(literature,
        'Garay et al. (2026) untersuchen eine multimodale Architektur mit Vibrations-, Akustik- '
        'und Thermografiesensorik. Die begutachtete Zeitschriftenfassung verwendet für die '
        'Vibrationserkennung zehn Zeitbereichsmerkmale aus 6-s-Fenstern eines Prüfstands mit '
        'gezielt erzeugter statischer Unwucht. Verglichen werden PCA, One-Class SVM, Isolation '
        'Forest und zwei Zahlenformate eines Autoencoders (Float32 und INT8). Neben einer '
        'Referenzplattform wird der Cortex-M4F des Sensorknotens vermessen. Beim Isolation '
        'Forest gehören Erkennungswerte zum Offline-Modell mit 100 Bäumen, die MCU-Laufzeit '
        'dagegen zu einer verkleinerten Implementierung mit sechs Bäumen. Qualitäts- und '
        'Laufzeitangaben müssen daher dem jeweiligen Modellartefakt zugeordnet werden. Andere '
        'Eingangsdaten, Fensterlängen und Zielplattformen verhindern eine unmittelbare '
        'Übertragung der Rangfolge auf den eigenen Lüfteraufbau.')
    for row in doc.tables[2].rows:
        if row.cells[0].text.startswith('Garay et al.'):
            change_text(row.cells[0].paragraphs[0], 'Garay et al. (2026)')
            change_text(row.cells[1].paragraphs[0],
                        'Statische Unwucht; zehn Zeitbereichsmerkmale; vier Verfahrensklassen mit zwei Autoencoder-Zahlenformaten')
            change_text(row.cells[3].paragraphs[0],
                        'Zielplattform, Datenrepräsentation und Modellartefakt zuordnen; keine direkte Rangfolge zum eigenen Aufbau')
    bibliography = find_paragraph(doc, 'Garay, C. E., Miranda Bonomi, F. A.')
    change_text(bibliography,
        'Garay, C. E., Miranda Bonomi, F. A., Mansilla, G. N., Fagre, M., Guzmán, S. G., '
        'Ritorto, P. A., Perez, F. I., & Katz, M. (2026). A Multimodal TinyML-Based Predictive '
        'Maintenance Architecture for Industrial IoT in the 6G Era. Sensors, 26(14), 4536. '
        'https://doi.org/10.3390/s26144536 (Abruf: 19.09.2026).')

    # Keep the existing pipeline drawing, but number all chapter-6 figures in reading order.
    pipeline = find_paragraph(doc, 'Abbildung 6-2: Datenfluss')
    change_text(pipeline, pipeline.text.replace('Abbildung 6-2:', 'Abbildung 6-4:', 1))
    training = find_paragraph(doc, 'Der Autoencoder wird mit Adam, Lernrate')
    change_text(training, training.text + ' Abbildung 6-2 zeigt den gespeicherten Lernverlauf.')
    quantile = find_paragraph(doc, 'Für jede Methode wird der Schwellenwert als 99. Perzentil')
    add_figure_before(quantile, 'training_history.png',
        'Abbildung 6-2: Trainings- und Validierungsverlust des Autoencoders. Gezeigt werden '
        'die protokollierten MSE-Werte der 100 Epochen, keine Erkennungsquoten. Der niedrigste '
        'Validierungsverlust liegt am Epochenlimit; vollständige Konvergenz ist damit nicht '
        'belegt. Quelle: B2, results/training_history.csv.',
        'MSE über 100 Epochen: Training endet bei 0,251450; Validierung bei 0,278935. '
        'Die beste Validierungsepoche ist 100, zugleich das Epochenlimit.')
    explanation = quantile.insert_paragraph_before(
        'Die Trainings- und Validierungsverluste nehmen im dokumentierten Verlauf ab. '
        'Dies zeigt die Anpassung an die normalen Aufnahmen, belegt aber weder eine '
        'allgemeine Erkennungsqualität noch eine optimale Trainingsdauer. Insbesondere '
        'ersetzt die Validierungskurve keinen unabhängigen Test mit Ereignisreferenzen.')
    explanation.paragraph_format.keep_together = True

    export = find_paragraph(doc, 'Die Autoencoder-Schwelle stammt aus der eingefrorenen Keras-Validierung.')
    calibration_intro = export.insert_paragraph_before(
        'Abbildung 6-3 zeigt die Verteilungen der 97 Kalibrierungsscores je Methode auf '
        'getrennten Scoreachsen. Die Autoencoder-Darstellung verwendet die ursprünglichen '
        'Keras-Rekonstruktionsfehler, aus denen die eingefrorene Schwelle stammt. Bei jeder '
        'Methode liegt eines dieser 97 Kalibrierungsfenster oberhalb der jeweiligen Schwelle. '
        'Das ist ein Befund auf den zur Grenzfestlegung verwendeten Daten, keine unabhängige '
        'Schätzung der Fehlalarmrate.')
    calibration_intro.paragraph_format.keep_with_next = True
    add_figure_before(export, 'calibration_scores.png',
        'Abbildung 6-3: Verteilung der Kalibrierungsscores und eingefrorene P99-Schwellen. '
        'Treppenlinien zeigen die empirischen Verteilungsfunktionen; senkrechte Linien '
        'kennzeichnen die separat linear interpolierten 99. Perzentile. Markiert sind '
        'jeweils die zwei größten Scores. Die Scoreachsen sind nicht untereinander '
        'kalibriert. Quelle: B2, results/validation_reconstruction_errors.csv (Keras); '
        'B1, bundle/validation_predictions.csv (Isolation Forest und RMS).',
        'Drei empirische Verteilungsfunktionen mit jeweils 97 Normalfenstern. '
        'P99-Schwellen: Autoencoder 0,477203, Isolation Forest 0,483355, RMS 1,296361. '
        'Je Methode liegt ein Kalibrierungswert darüber; keine Test-Fehlalarmraten.')

    null_text = find_paragraph(doc, 'cues.jsonl speichert die Bedienereignisse;')
    old = ('Eine Null oder ein fehlender Wert für true_detection_delay_ms, detection_rate '
           'beziehungsweise false_alarm_rate bedeutet nicht null Millisekunden oder null '
           'Prozent, sondern einen nicht bestimmbaren Nachweis.')
    assert old in null_text.text
    change_text(null_text, null_text.text.replace(old,
        'Der JSON-Wert null oder ein fehlender Eintrag für true_detection_delay_ms, '
        'detection_rate beziehungsweise false_alarm_rate kennzeichnet einen nicht '
        'bestimmbaren Wert; daraus dürfen nicht 0 ms oder 0 % abgeleitet werden. '
        'Die Zahl 0 ist dagegen ein numerischer Wert und kein Platzhalter für fehlende Daten.'))

    # Sort the appended abbreviations without changing any definitions or cell formatting.
    abbreviation_table = doc.tables[0]
    for row in sorted(list(abbreviation_table.rows)[1:], key=lambda r: r.cells[0].text.casefold()):
        abbreviation_table._tbl.append(row._tr)
    doc.core_properties.modified = datetime.now(timezone.utc)
    doc.core_properties.comments = (
        'Explorative Smartphone-Vibrationsstudie; Trainings- und Kalibrierungsdiagramme '
        'aus gespeicherten Originaldaten ergänzt. Keine Änderung von Messdaten oder Modellen.')
    assert len(doc.tables) == original_tables
    assert len(doc.inline_shapes) == original_shapes + 2
    assert 'Abgabedatum: 27.10.2026' in '\n'.join(p.text for p in doc.paragraphs)
    staged = OUT / 'Masterarbeit_ueberarbeitet_staged.docx'
    doc.save(staged)
    return staged


def main():
    assert digest(SOURCE) == EXPECTED_SOURCE_HASH, 'Source changed since review; inspect before editing.'
    OUT.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)
    assert digest(BACKUP) == EXPECTED_SOURCE_HASH
    hashes = {key: digest(path) for key, path in SOURCES.items()}
    metrics = make_figures()
    staged = revise_document()
    assert hashes == {key: digest(path) for key, path in SOURCES.items()}
    assert digest(SOURCE) == EXPECTED_SOURCE_HASH
    manifest = {
        'source': str(SOURCE.relative_to(ROOT)), 'source_sha256': EXPECTED_SOURCE_HASH,
        'backup': str(BACKUP.relative_to(ROOT)), 'staged': str(staged.relative_to(ROOT)),
        'data_sources': {key: {'path': str(path.relative_to(ROOT)), 'sha256': hashes[key]}
                         for key, path in SOURCES.items()},
        'metrics': metrics,
        'source_update': {
            'reference': 'Garay et al. (2026), Sensors 26(14), 4536',
            'doi': 'https://doi.org/10.3390/s26144536',
            'primary_institutional_record': 'https://oulurepo.oulu.fi/handle/10024/64560',
            'verified': ['bibliography', 'modalities', '6-second windows and 10 features',
                         'four algorithms / five variants', '100-tree offline vs 6-tree MCU IF'],
            'checked_on': '2026-09-19',
        },
        'changes': ['two evidence-based figures with explanations',
                    'pipeline figure renumbered 6-2 to 6-4',
                    'Garay journal version and artifact distinction',
                    'JSON null versus numeric zero clarified', 'abbreviations alphabetized'],
        'preserved': ['experiment data', 'model weights', 'thresholds', 'resource measurements',
                      'pilot-study limitations', 'cover dates pending user confirmation'],
    }
    (OUT / 'revision_manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False)+'\n')
    print(json.dumps({'backup': str(BACKUP), 'staged': str(staged), 'metrics': metrics}, indent=2))


if __name__ == '__main__':
    main()
