#!/usr/bin/env python3
"""Build a staged Word revision from a preserved source; never alter measurements."""
from pathlib import Path
from copy import deepcopy
import hashlib
import importlib.util
import json
import re
import shutil
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/phone_vibration_word_revision_20260919'
TARGET=ROOT/'docs/Akz_Masterarbeit_Bericht(3).docx'
BACKUP=ROOT/'docs/backups/Akz_Masterarbeit_Bericht(3)_vor_handyversuch_20260919.docx'


def sha(path):
    with path.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    BACKUP.parent.mkdir(parents=True,exist_ok=True)
    if not BACKUP.exists():shutil.copy2(TARGET,BACKUP)
    spec=importlib.util.spec_from_file_location('thesis_helpers',ROOT/'results/upright_v4_finalization_20260912_103145/build_thesis.py')
    h=importlib.util.module_from_spec(spec);spec.loader.exec_module(h)
    doc=Document(BACKUP);body=doc._element.body;blocks=list(body)
    assert Paragraph(blocks[230],doc._body).text=='Implementierung'
    assert Paragraph(blocks[449],doc._body).text=='Literaturverzeichnis'
    bibliography_xml=b''.join(x.xml.encode() for x in blocks[449:])
    changes=[]
    def replace(index,text):
        p=Paragraph(blocks[index],doc._body);old=p.text
        p.clear();h.text_runs(p,text);changes.append({'body_index':index,'old':old,'new':text})

    replace(29,'Diese Arbeit vergleicht RMS-Schwelle, Isolation Forest und TFLite-Autoencoder zur lokalen Anomalieerkennung an einem Lüfteraufbau. Die Hardware bleibt ein Raspberry Pi 5 mit ADXL345 und einem extern versorgten Lüfter bei 75 % PWM. Im aktuellen Versuch wird eine äußere Anregung durch ein auf dem Tisch vibrierendes Handy erzeugt. Alle Verfahren verarbeiten nicht überlappende Fenster mit 128 XYZ-Punkten. Modelle, Skalierung und P99-Schwellen des am 18. September 2026 erstellten Normalprofils bleiben während der Handyversuche unverändert.')
    replace(30,'Am 19. September 2026 wurden drei ausgewählte Einzelaufnahmen und ein gemeinsamer Live-Lauf ausgewertet. Im gemeinsamen Lauf erhalten die drei Verfahren dieselben 197 Rohfenster. Der Autoencoder gibt den ersten Alarm aus; Isolation Forest folgt 19,061 ms später im selben Fenster 108. RMS überschreitet seine Schwelle erst in Fenster 109 und gibt den ersten Alarm 615,659 ms nach dem Autoencoder aus. Damit zeigt der Autoencoder in diesem Lauf den günstigsten beobachteten Alarmausgabezeitpunkt. Im Replay der drei Einzelaufnahmen liegt er einmal ein Fenster vor beiden Baselines und zweimal im gleichen ersten Alarmfenster. Diese Ergebnisse beschreiben wenige Ereignisse und keine allgemeine Überlegenheit.')
    replace(31,'Der gemeinsame Prozess benötigt einen CPU-Median von 13,04 % eines logischen Kerns und maximal 174,42 MiB abgetasteten RSS. RMS hat den geringsten, Isolation Forest den höchsten Verarbeitungsaufwand. Elektrische Leistung und Energie wurden nicht gemessen. Die Aufforderung zum Anrufen und die Rückmeldung zum Ende sind keine physikalischen Zeitreferenzen; eine absolute Erkennungsverzögerung, F1 oder bestätigte Fehlalarmrate lassen sich aus der Handyreihe daher nicht bestimmen. Frühere Versuche mit einem anderen Modell- und Vorverarbeitungsstand bleiben als getrennte Prüfreihe erhalten, einschließlich ihrer negativen Erkennungsergebnisse. Der Bericht ist eine nachvollziehbare Pilotstudie mit dokumentierten Aussagegrenzen.')
    replace(50,'Der Beitrag ist ein Methodenvergleich auf identischen Sensoreingaben mit nachvollziehbarer Schwellenkalibrierung und Messung des technischen Aufwands. Der aktuelle Schwerpunkt liegt auf der Reihenfolge der Alarmausgaben bei einer Handyvibration. Frühere Prüfungen zur Erkennungsqualität und zu geänderten Normalbedingungen werden getrennt eingeordnet; sie werden durch den neuen Versuch nicht nachträglich umgedeutet.')
    replace(54,'FF2: Wie unterscheiden sich Hauptverfahren und Baselines bei gleichen Testfenstern und einer einheitlichen Schwellenwertkalibrierung auf Normaldaten hinsichtlich Precision, Recall, F1-Score und Fehlalarmrate? Ergänzend wird in der Handyreihe deskriptiv untersucht, welche Methode zuerst ein Alarmfenster markiert und ihre Alarmausgabe bereitstellt. Diese ergänzende Zeitfrage ersetzt die ursprüngliche Qualitätsfrage nicht.')
    replace(57,'Die ursprünglichen Arbeitshypothesen bleiben als Bewertungsmaßstab erhalten. Der Bericht unterscheidet die frühere Prüfreihe mit definierten Zustandslabels und die nachfolgende explorative Handyreihe mit Bedienmarkern. Training und Schwellenkalibrierung der Handyreihe erfolgen auf vorher aufgenommenen Normaldaten; an den Störungsaufnahmen werden keine Schwellen angepasst. Wegen fehlender unabhängiger Zeitreferenz und weniger Ereignisse liefert die neue Reihe keinen zusätzlichen vollständigen Test von H1 oder H3. Die ergänzende Rangfolge der Alarmausgaben ist ein deskriptiver Befund, keine rückwirkend vorab formulierte Hypothese.')
    replace(64,'Die überarbeitete Fassung berücksichtigt die Handyversuche vom 19. September 2026 und das zugehörige Normalprofil vom 18. September. Die Hardwareauswahl bleibt unverändert. Frühere Messungen vom 11./12. September mit separater Platte und einem anderen Modellstand werden als historische Prüfreihe kenntlich gemacht und im Anhang nachgewiesen. Die Handyreihe verwendet keine Platte. Nicht vorhandene Wiederholungen, Drehzahlmessungen, physische Vibrationszeitreferenzen und elektrische Verbrauchsmessungen werden nicht als erbracht behandelt. Teilnachweise der Messqualität und Wiederholbarkeit werden nicht nachträglich zu vollständiger Erfüllung erklärt.')
    replace(66,'Kapitel 2 und 3 erläutern Grundlagen und Vorarbeiten. Kapitel 4 und 5 begründen Anforderungen und Hardwareauswahl. Kapitel 6 beschreibt den aktuellen Handyaufbau, die Implementierung und das Protokoll. Kapitel 7 trennt Einzelbetrieb, Replay und gemeinsamen Live-Vergleich einschließlich Raspberry-Pi-Ressourcen. Diskussion, Ausblick und Fazit ordnen die Ergebnisse ein. Der Anhang erhält die früheren Prüfergebnisse und verknüpft alle digitalen Nachweise.')
    replace(107,'P95 und P99 beschreiben hohe empirische Perzentile der gemessenen Latenzen. Sie werden zusammen mit Fristüberschreitungen und Erfassungspausen ausgewertet. Davon zu unterscheiden sind das erste als auffällig markierte Fenster und der Zeitabstand zwischen den ersten Alarmausgaben mehrerer Methoden. Eine Verzögerung ab physischem Störungsbeginn setzt eine unabhängige Zeitreferenz voraus. Trainierte Modelle können für die lokale Ausführung in ein anderes Format konvertiert werden (Google, 2026a). Danach müssen Scores und Entscheidungen erneut geprüft werden.')
    replace(127,'Dafür werden gleiche Testfenster und eine gemeinsame Kalibrierungsregel für FF2 sowie Messungen am tatsächlich ausgeführten Modell für FF3 verbunden. Die frühere Untersuchung geänderter Normalbedingungen adressiert FF4 separat. Die ergänzende Handyreihe vergleicht zuerst Einzelprozesse und anschließend dieselben Rohfenster im Replay und im gemeinsamen Live-Prozess. So wird zwischen unterschiedlichen physikalischen Anregungen, identischen Eingaben und zusätzlicher Konkurrenz um Rechenzeit unterschieden.')
    changes.append({'body_index':128,'old':Paragraph(blocks[128],doc._body).text,'new':'','reason':'Redundanter Kapitelübergang entfernt, der nach der Aktualisierung allein auf einer Seite stünde.'})
    body.remove(blocks[128])
    replace(175,'Als konkrete Variante wurde der ARCTIC P12 Pro PST vorgesehen, dessen Datenblatt PWM-Ansteuerung und ein Drehzahlsignal ausweist (ARCTIC, o. J.). Ein PWM-Tastgrad ist eine Stellgröße und kein Nachweis einer konstanten Drehzahl. Die Konzeption sah deshalb die Erfassung der tatsächlichen Drehzahl vor. Dieser Nachweis wurde bisher nicht erbracht; die Betriebspunkte werden durch ihre dokumentierten PWM-Vorgaben beschrieben. Für FF4 wurde in der früheren Prüfreihe ein zweiter normaler Betriebspunkt bei unveränderter Montage untersucht; diese Werte bleiben im Anhang erhalten. Den aktuellen Handyaufbau und Messablauf beschreibt Kapitel 6, die zugehörigen Befunde und Grenzen Kapitel 7.')
    # The roadmap is already in Chapter 1. Remove a redundant transition that
    # would otherwise occupy a page of its own before the Chapter 6 page break.
    changes.append({'body_index':229,'old':Paragraph(blocks[229],doc._body).text,'new':'','reason':'Redundanter Übergang entfernt; kein verwaister Absatz auf eigener Seite.'})
    body.remove(blocks[229])
    for e in blocks[:230]:
        if e.tag==qn('w:p'):
            p=Paragraph(e,doc._body)
            if 'Laborversuch' in p.text:
                old=p.text;new=old.replace('Laborversuch','Versuchsaufbau');p.clear();p.add_run(new);changes.append({'old':old,'new':new})

    clones={
        'ORIGINAL_ARCHITECTURE':[deepcopy(blocks[266])],
        'HISTORICAL_LATENCY':[deepcopy(blocks[359])],
        'HISTORICAL_RESOURCE':[deepcopy(blocks[366])],
        'EVIDENCE_INDEX':[deepcopy(blocks[424])],
        'HISTORICAL_COUNTS':[deepcopy(blocks[431])],
        'HISTORICAL_METRICS':[deepcopy(blocks[434])],
        'HISTORICAL_PER_RECORD':[deepcopy(blocks[440])],
    }
    anchor=blocks[449]
    for e in blocks[230:449]:body.remove(e)
    # Keep front matter, hardware-selection tables, equations, bibliography and sections.
    # Only the new main text is inserted before the original bibliography.
    bookmark_id=10000
    def put(element):
        parent=element.getparent()
        if parent is not None:parent.remove(element)
        anchor.addprevious(element)
    def paragraph(text='',style=None):
        p=doc.add_paragraph(style=style);h.text_runs(p,text);put(p._p);return p
    def mark(p,prefix):
        nonlocal bookmark_id
        h.bookmark(p,f'phone_{prefix}_{bookmark_id}',bookmark_id);bookmark_id+=1
    def keep_table(t):
        # The new tables fit on one page; avoid a stranded last row on the next.
        for row_index,row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next=(row_index<len(t.rows)-1)
    lines=(OUT/'manuscript_update.md').read_text().splitlines();i=0;appendix=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line[level:].strip()
            if title=='Anhang':appendix=True
            p=paragraph(title,f'Heading {level}');h.set_num(p,level-1,0 if appendix else 2)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT;mark(p,'heading');continue
        if line.startswith('{{'):
            key=line[2:-2]
            for original in clones[key]:
                e=deepcopy(original);put(e)
                if key=='EVIDENCE_INDEX':
                    t=Table(e,doc._body)
                    for cells in [('B12','Handyversuche, eingefrorenes Bündel, Einzel-/Replay-/Live-Vergleich und Präsentation','results/phone_vibration_pilot_20260919/'),('B13','Word-Revision, neue Diagramme, Quellenhashes und Layoutprüfung','results/phone_vibration_word_revision_20260919/')]:
                        row=t.add_row()
                        for c,v in zip(row.cells,cells):c.text=v
                    h.style_table(t,is_appendix=True)
                    for row in t.rows:
                        for p in row.cells[2].paragraphs:
                            for run in p.runs:run.text=run.text.replace('/','/\u200b').replace('_','_\u200b')
                if e.tag==qn('w:tbl'):keep_table(Table(e,doc._body))
            continue
        if line.startswith('!['):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line);assert match,line
            caption,path=match.groups();p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(0)
            p.add_run().add_picture(str(OUT/path),width=Cm(16))
            cp=paragraph(caption,'Abbildungsbeschriftung');mark(cp,'figure');continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            data=[[v.strip() for v in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
            t=doc.add_table(rows=len(data),cols=len(data[0]));put(t._tbl)
            for row,vals in zip(t.rows,data):
                assert len(vals)==len(data[0])
                for c,v in zip(row.cells,vals):c.text=v
            h.style_table(t);keep_table(t);continue
        if line.startswith('Tabelle '):
            p=paragraph(line,'Tabellenbeschriftung');mark(p,'table');continue
        paragraph(line)
    assert b''.join(x.xml.encode() for x in list(body)[list(body).index(anchor):])==bibliography_xml
    # Heading numbering remains tied to the existing main list; appendices are unnumbered.
    doc.core_properties.modified=__import__('datetime').datetime.now(__import__('datetime').timezone.utc)
    doc.core_properties.comments='Revision: Handyversuch 19.09.2026; archivierte Quellen unverändert.'
    stage=OUT/'Akz_Masterarbeit_Bericht(3)_staged.docx';doc.save(stage)
    audit={'source':str(BACKUP.relative_to(ROOT)),'source_sha256':sha(BACKUP),'staged':str(stage.relative_to(ROOT)),'staged_sha256':sha(stage),'retained_bibliography_xml_sha256':hashlib.sha256(bibliography_xml).hexdigest(),'preserved':'Deckblatt, Erklärungen, Grundlagen, Hardwareauswahl, Gleichungen, Literatur, Abschnittseinstellungen; historische Zahlentabellen unverändert','target_replaced':False,'paragraph_changes_before_chapter6':changes,'manuscript_sha256':sha(OUT/'manuscript_update.md'),'figures':json.loads((OUT/'figure_manifest.json').read_text())}
    (OUT/'revision_manifest.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'output':str(stage),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'inline_images':len(doc.inline_shapes)},indent=2))


if __name__=='__main__':main()
